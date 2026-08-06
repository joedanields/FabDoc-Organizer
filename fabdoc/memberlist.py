"""Import the member list exported from the structural model.

Tekla, SDS2 and friends all export something different - .xlsx, .csv, or a plain
text dump. This module reads all three, guesses which column holds the member
mark, and reports what it guessed so the engineer can override it.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

# Header text that signals the member-mark column, best candidate first.
_MEMBER_HEADER_HINTS = [
    "assembly mark", "assembly position", "piece mark", "member name",
    "member mark", "part mark", "assembly", "mark", "member", "part",
    "position", "name", "pos",
]

_SKIP_VALUES = {"", "-", "n/a", "na", "none", "null"}


@dataclass
class MemberList:
    """Member marks read from a model export."""

    members: list[str] = field(default_factory=list)
    source: str = ""
    column_name: str = ""
    column_index: int = 0
    headers: list[str] = field(default_factory=list)
    rows_read: int = 0

    @property
    def count(self) -> int:
        return len(self.members)


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _pick_column(headers: list[str]) -> int:
    """Index of the column most likely to hold member marks."""
    lowered = [h.strip().lower() for h in headers]
    for hint in _MEMBER_HEADER_HINTS:
        for idx, head in enumerate(lowered):
            if head == hint:
                return idx
    for hint in _MEMBER_HEADER_HINTS:
        for idx, head in enumerate(lowered):
            if hint in head:
                return idx
    return 0


def _looks_like_header(row: list[str]) -> bool:
    """A row is a header if it has text and no cell is purely numeric."""
    filled = [c for c in row if c.strip()]
    if not filled:
        return False
    return not any(c.strip().replace(".", "", 1).isdigit() for c in filled)


def read_excel(path: Path, column: int | str | None = None,
               sheet: str | None = None) -> MemberList:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    return _from_rows(rows, source=str(path), column=column)


def read_csv(path: Path, column: int | str | None = None) -> MemberList:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [[_clean(c) for c in row] for row in csv.reader(fh, dialect)]
    return _from_rows(rows, source=str(path), column=column)


def read_text(path: Path) -> MemberList:
    """One member per line, for plain dumps with no structure."""
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        lines = [line.strip() for line in fh]
    members = [ln for ln in lines if ln and ln.lower() not in _SKIP_VALUES]
    return MemberList(members=members, source=str(path),
                      column_name="(one per line)", rows_read=len(lines))


def _from_rows(rows: list[list[str]], source: str,
               column: int | str | None = None) -> MemberList:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return MemberList(source=source)

    header_row: list[str] = []
    body = rows
    if _looks_like_header(rows[0]):
        header_row = rows[0]
        body = rows[1:]

    if isinstance(column, int):
        index = column
    elif isinstance(column, str) and column:
        lowered = [h.strip().lower() for h in header_row]
        index = lowered.index(column.strip().lower()) if column.strip().lower() in lowered else 0
    else:
        index = _pick_column(header_row) if header_row else 0

    members: list[str] = []
    for row in body:
        if index < len(row):
            value = row[index].strip()
            if value and value.lower() not in _SKIP_VALUES:
                members.append(value)

    return MemberList(
        members=members,
        source=source,
        column_name=header_row[index] if header_row and index < len(header_row) else f"Column {index + 1}",
        column_index=index,
        headers=header_row,
        rows_read=len(body),
    )


def read_member_list(path: str | Path, column: int | str | None = None,
                     sheet: str | None = None) -> MemberList:
    """Read a member list from .xlsx/.xlsm, .csv/.tsv, or .txt."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_excel(p, column=column, sheet=sheet)
    if suffix in (".csv", ".tsv"):
        return read_csv(p, column=column)
    if suffix in (".txt", ".lst", ".dat", ""):
        return read_text(p)
    if suffix == ".xls":
        raise ValueError(
            "Legacy .xls is not supported. Re-save the export as .xlsx or .csv."
        )
    raise ValueError(f"Unsupported member list format: {suffix}")


def preview_columns(path: str | Path, sheet: str | None = None) -> list[str]:
    """Header names available in a member list, for the column picker."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
            for row in ws.iter_rows(values_only=True):
                cells = [_clean(c) for c in row]
                if any(cells):
                    return cells if _looks_like_header(cells) else []
        finally:
            wb.close()
        return []
    if suffix in (".csv", ".tsv"):
        result = read_csv(p)
        return result.headers
    return []


def sheet_names(path: str | Path) -> list[str]:
    """Worksheet names in an Excel member list."""
    p = Path(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
