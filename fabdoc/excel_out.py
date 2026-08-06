"""Write the drawing register and validation report to Excel workbooks."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .categories import safe_sheet_name
from .register import Register, CategoryRegister

# --- Shared styling ---------------------------------------------------------

_TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_LABEL_FONT = Font(bold=True, size=10, color="1F4E79")
_VALUE_FONT = Font(size=10)
_HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")   # amber: check this row
_ERROR_FILL = PatternFill("solid", fgColor="F8CBAD")    # orange: file failed
_OK_FILL = PatternFill("solid", fgColor="E2EFDA")
_MISSING_FILL = PatternFill("solid", fgColor="F8CBAD")
_EXTRA_FILL = PatternFill("solid", fgColor="FFF2CC")

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _style_header_band(ws: Worksheet, register: Register, category: CategoryRegister,
                       width: int) -> int:
    """Write the project metadata band. Returns the next free row."""
    last_col = get_column_letter(width)
    meta = register.meta

    ws.merge_cells(f"A1:{last_col}1")
    cell = ws["A1"]
    cell.value = "DRAWING REGISTER"
    cell.font = _TITLE_FONT
    cell.fill = _TITLE_FILL
    cell.alignment = _CENTER
    ws.row_dimensions[1].height = 22

    pairs = [
        ("Project Title", meta.title),
        ("Issue Date", meta.date_display),
        ("Zone", meta.zone),
        ("Package", meta.package),
        ("Category", category.name),
        ("Drawings", str(category.total)),
        ("Source Folder", str(category.folder)),
        ("Generated", datetime.now().strftime("%d-%b-%Y %H:%M")),
    ]

    row = 2
    col = 1
    for label, value in pairs:
        if not value:
            continue
        lc = ws.cell(row=row, column=col, value=f"{label}:")
        lc.font = _LABEL_FONT
        lc.alignment = _LEFT
        vc = ws.cell(row=row, column=col + 1, value=value)
        vc.font = _VALUE_FONT
        vc.alignment = _LEFT
        col += 2
        if col > width - 1:
            col = 1
            row += 1
    return row + 2 if col == 1 else row + 2


def _write_category_sheet(ws: Worksheet, register: Register,
                          category: CategoryRegister, include_source: bool) -> None:
    # Columns are exactly those named in the requirement; Title and Date repeat
    # per row so registers from several issues can be stacked and filtered.
    headers = ["Title", "Date", "S.No", "Member Name", "Revision No"]
    if include_source:
        headers += ["Source File", "Notes"]

    header_row = _style_header_band(ws, register, category, len(headers))

    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[header_row].height = 18

    title = register.meta.title
    date_text = register.meta.date_display

    row = header_row + 1
    for rec in category.records:
        seq: object = int(rec.seq_no) if rec.seq_no.isdigit() else rec.seq_no
        values: list[object] = [title, date_text, seq, rec.member_name, rec.revision]
        if include_source:
            values += [rec.source_file, rec.error or rec.note_text]

        for idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=idx, value=value)
            cell.border = _BORDER
            cell.alignment = _CENTER if idx in (2, 3, 5) else _LEFT
            if rec.error:
                cell.fill = _ERROR_FILL
            elif rec.needs_review:
                cell.fill = _REVIEW_FILL
        row += 1

    # Freeze the header and enable filtering: registers run to thousands of rows.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if row > header_row + 1:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{row - 1}"

    widths = [34, 14, 8, 24, 12, 40, 34]
    for idx, width in enumerate(widths[: len(headers)], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_summary_sheet(ws: Worksheet, register: Register) -> None:
    """Counts per category, so a large register can be sanity-checked at a glance."""
    meta = register.meta
    ws.merge_cells("A1:D1")
    cell = ws["A1"]
    cell.value = "REGISTER SUMMARY"
    cell.font = _TITLE_FONT
    cell.fill = _TITLE_FILL
    cell.alignment = _CENTER

    row = 3
    for label, value in (
        ("Project Title", meta.title),
        ("Issue Date", meta.date_display),
        ("Zone", meta.zone),
        ("Package", meta.package),
        ("Revision", meta.revision),
        ("Source Folder", str(register.project_folder)),
        ("Folder Name", meta.folder_name),
        ("Generated", datetime.now().strftime("%d-%b-%Y %H:%M")),
    ):
        if not value:
            continue
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value).font = _VALUE_FONT
        row += 1

    row += 1
    for idx, name in enumerate(["Category", "Drawings", "Needs Review", "Errors"], start=1):
        c = ws.cell(row=row, column=idx, value=name)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
    row += 1

    for cat in register.categories:
        for idx, value in enumerate(
            [cat.name, cat.total, cat.review_count, cat.error_count], start=1
        ):
            c = ws.cell(row=row, column=idx, value=value)
            c.border = _BORDER
            c.alignment = _CENTER if idx > 1 else _LEFT
        row += 1

    for idx, value in enumerate(
        ["TOTAL", register.total, register.review_count,
         sum(c.error_count for c in register.categories)], start=1
    ):
        c = ws.cell(row=row, column=idx, value=value)
        c.font = Font(bold=True)
        c.border = _BORDER
        c.alignment = _CENTER if idx > 1 else _LEFT

    row += 2
    note = ws.cell(
        row=row, column=1,
        value="Amber rows in the register were extracted from the filename or "
              "not found at all - confirm them against the drawing. Orange rows "
              "are files that could not be read.",
    )
    note.font = Font(italic=True, size=9, color="808080")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

    for idx, width in enumerate([26, 40, 16, 12], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_register(register: Register, output_path: str | Path,
                   include_source: bool = True, include_summary: bool = True) -> Path:
    """Write the register to an .xlsx workbook, one worksheet per category."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    if include_summary:
        _write_summary_sheet(wb.create_sheet("Summary"), register)

    used: set[str] = {"Summary"} if include_summary else set()
    for cat in register.categories:
        ws = wb.create_sheet(safe_sheet_name(cat.name, used))
        _write_category_sheet(ws, register, cat, include_source)

    if not register.categories and not include_summary:
        wb.create_sheet("Register")

    wb.save(out)
    return out


def suggest_register_name(register: Register) -> str:
    """A filename for the register based on project metadata."""
    meta = register.meta
    bits = [b for b in (meta.title, meta.zone, meta.date_display) if b]
    stem = " - ".join(bits) if bits else "Drawing Register"
    safe = "".join(ch for ch in stem if ch not in '<>:"/\\|?*').strip()
    return f"{safe[:120]} - Drawing Register.xlsx"


# ---------------------------------------------------------------------------
# Validation report
# ---------------------------------------------------------------------------


def write_validation_report(result, output_path: str | Path) -> Path:
    """Write a ValidationResult to a multi-sheet workbook."""
    from .validate import ValidationResult  # local import avoids a cycle

    assert isinstance(result, ValidationResult)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.merge_cells("A1:D1")
    cell = ws["A1"]
    cell.value = "MEMBER VALIDATION REPORT"
    cell.font = _TITLE_FONT
    cell.fill = _TITLE_FILL
    cell.alignment = _CENTER

    row = 3
    for label, value in (
        ("Project", result.project_title),
        ("Issue Date", result.project_date),
        ("Drawing Register", result.register_source),
        ("Model Member List", result.model_source),
        ("Generated", datetime.now().strftime("%d-%b-%Y %H:%M")),
    ):
        if not value:
            continue
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value).font = _VALUE_FONT
        row += 1

    row += 1
    counts = [
        ("Members in model", result.model_count, None),
        ("Members in drawings", result.drawing_count, None),
        ("Matched", len(result.matched), _OK_FILL),
        ("Missing in drawings", len(result.missing_in_drawings), _MISSING_FILL),
        ("Not in model", len(result.extra_in_drawings), _EXTRA_FILL),
        ("Duplicated in drawings", len(result.duplicates_in_drawings), _EXTRA_FILL),
    ]
    for idx, name in enumerate(["Result", "Count"], start=1):
        c = ws.cell(row=row, column=idx, value=name)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _BORDER
    row += 1
    for label, value, fill in counts:
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=value)
        for c in (lc, vc):
            c.border = _BORDER
            if fill:
                c.fill = fill
        vc.alignment = _CENTER
        row += 1

    row += 1
    verdict = ws.cell(row=row, column=1, value=result.verdict)
    verdict.font = Font(bold=True, size=12,
                        color="375623" if result.is_clean else "C00000")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

    for idx, width in enumerate([30, 44, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    def detail_sheet(title: str, headers: list[str], rows: list[list[object]],
                     fill: PatternFill | None) -> None:
        sheet = wb.create_sheet(safe_sheet_name(title))
        for idx, name in enumerate(headers, start=1):
            c = sheet.cell(row=1, column=idx, value=name)
            c.font = _HEADER_FONT
            c.fill = _HEADER_FILL
            c.alignment = _CENTER
            c.border = _BORDER
        for r_idx, values in enumerate(rows, start=2):
            for c_idx, value in enumerate(values, start=1):
                c = sheet.cell(row=r_idx, column=c_idx, value=value)
                c.border = _BORDER
                if fill:
                    c.fill = fill
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        for idx in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(idx)].width = 26

    detail_sheet(
        "Missing in Drawings", ["Member Name (model)"],
        [[m] for m in result.missing_in_drawings], _MISSING_FILL,
    )
    detail_sheet(
        "Not in Model", ["Member Name (drawing)", "Category", "S.No", "Revision", "Source File"],
        [[m.member_name, m.category, m.seq_no, m.revision, m.source_file]
         for m in result.extra_records], _EXTRA_FILL,
    )
    detail_sheet(
        "Matched", ["Member Name", "Category", "S.No", "Revision", "Source File"],
        [[m.member_name, m.category, m.seq_no, m.revision, m.source_file]
         for m in result.matched_records], _OK_FILL,
    )
    if result.duplicates_in_drawings:
        detail_sheet(
            "Duplicates", ["Member Name", "Times in Drawings", "Source Files"],
            [[name, len(files), ", ".join(files)]
             for name, files in result.duplicates_in_drawings.items()], _EXTRA_FILL,
        )

    wb.save(out)
    return out
