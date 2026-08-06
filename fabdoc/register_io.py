"""Read a previously generated register workbook back into a Register.

Lets an engineer validate against a register produced last week without
re-scanning several thousand PDFs.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .extract import DrawingRecord
from .folder_meta import ProjectMeta
from .register import CategoryRegister, Register, sort_records

_REQUIRED = {"member name"}


def _find_header_row(rows: list[list[str]]) -> int:
    """Row index holding the table headers, or -1 when absent."""
    for idx, row in enumerate(rows[:40]):
        lowered = {str(c).strip().lower() for c in row if c}
        if _REQUIRED <= lowered:
            return idx
    return -1


def read_register(path: str | Path) -> Register:
    """Load a register workbook written by :func:`fabdoc.excel_out.write_register`."""
    p = Path(path)
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        meta = ProjectMeta(folder_name=p.stem, title=p.stem)
        register = Register(meta=meta, project_folder=p.parent)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = [["" if c is None else str(c).strip() for c in row]
                    for row in ws.iter_rows(values_only=True)]
            if not rows:
                continue

            if sheet_name.lower() == "summary":
                for row in rows:
                    if len(row) >= 2 and row[0]:
                        label = row[0].strip().rstrip(":").lower()
                        if label == "project title" and row[1]:
                            meta.title = row[1]
                        elif label == "issue date" and row[1]:
                            meta.date_text = row[1]
                        elif label == "zone" and row[1]:
                            meta.zone = row[1]
                        elif label == "package" and row[1]:
                            meta.package = row[1]
                continue

            header_idx = _find_header_row(rows)
            if header_idx < 0:
                continue
            headers = [str(c).strip().lower() for c in rows[header_idx]]

            def col(*names: str) -> int:
                for name in names:
                    if name in headers:
                        return headers.index(name)
                return -1

            i_seq = col("s.no", "sno", "s no", "seq", "sequence")
            i_member = col("member name", "member", "mark")
            i_rev = col("revision no", "revision", "rev")
            i_src = col("source file", "file")
            i_title = col("title")
            i_date = col("date")

            records: list[DrawingRecord] = []
            for row in rows[header_idx + 1:]:
                if i_member < 0 or i_member >= len(row):
                    continue
                member = row[i_member].strip()
                if not member:
                    continue
                rec = DrawingRecord(category=sheet_name, member_name=member)
                if 0 <= i_seq < len(row):
                    rec.seq_no = row[i_seq].strip()
                if 0 <= i_rev < len(row):
                    rec.revision = row[i_rev].strip()
                if 0 <= i_src < len(row):
                    rec.source_file = row[i_src].strip()
                records.append(rec)

                if not meta.title and 0 <= i_title < len(row):
                    meta.title = row[i_title].strip()
                if not meta.date_text and 0 <= i_date < len(row):
                    meta.date_text = row[i_date].strip()

            if records:
                register.categories.append(
                    CategoryRegister(name=sheet_name, folder=p.parent,
                                     records=sort_records(records))
                )
        return register
    finally:
        wb.close()
