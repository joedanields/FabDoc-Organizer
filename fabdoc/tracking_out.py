"""Write the cumulative package tracker to an Excel workbook.

The per-issue register answers "what is in this delivery". This answers "what has
happened to this package", which is the question asked in a progress meeting:
which revision is each member on, what shipped when, and what is still held up.

Kept out of ``excel_out`` because it renders a different object (a chain, not a
register) and that module is already the larger of the two.
"""

from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .categories import safe_sheet_name
from .excel_out import (_BORDER, _CENTER, _HEADER_FILL, _HEADER_FONT, _LABEL_FONT,
                        _LEFT, _OK_FILL, _REVIEW_FILL, _TITLE_FILL, _TITLE_FONT,
                        _VALUE_FONT, _save_workbook)
from . import sequencing
from .tracking import (STAGE_IFA, STAGE_IFF, PackageChain, split_member_id)

_IFA_FILL = PatternFill("solid", fgColor="DDEBF7")    # blue: still in approval
_IFF_FILL = PatternFill("solid", fgColor="E2EFDA")    # green: released to shop
_HOLD_FILL = PatternFill("solid", fgColor="F8CBAD")   # orange: not shipped
_NEW_FILL = PatternFill("solid", fgColor="FFF2CC")    # amber: new this issue
# "H" carries the on-hold state in text as well as in colour.
_HOLD_FONT = Font(bold=True, size=10, color="843C0C")
_TOTAL_FONT = Font(bold=True, size=10)
# The band heading inside a history sheet, matching the register.
# Grey, not a warning colour: a dropped member is out of scope, not a problem.
_DROP_FILL = PatternFill("solid", fgColor="D9D9D9")
_DROP_FONT = Font(bold=True, size=10, color="595959")
_BAND_FILL = PatternFill("solid", fgColor="C5E0B4")
_BAND_FONT = Font(bold=True, size=10, color="375623")
# Red, and used nowhere else on the sheet: a revision that skipped a rung is the
# one mark on the grid that says the tracker itself is missing something.
_REV_ERR_FILL = PatternFill("solid", fgColor="FFC7CE")
_REV_ERR_FONT = Font(bold=True, size=10, color="9C0006")


def _title(ws: Worksheet, text: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font = _TITLE_FONT
    cell.fill = _TITLE_FILL
    cell.alignment = _CENTER
    ws.row_dimensions[1].height = 22


def _headers(ws: Worksheet, names: list[str], row: int) -> None:
    for idx, name in enumerate(names, start=1):
        cell = ws.cell(row=row, column=idx, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.row_dimensions[row].height = 18


def _category_counts(step, chain: PackageChain, label: str) -> "OrderedDict[str, dict]":
    """One issue's numbers, split by drawing category.

    Every list on a step holds member identities, and an identity carries its
    category, so the split is a regroup rather than a second pass over the data.
    """
    counts: "OrderedDict[str, dict]" = OrderedDict()

    def bucket(name: str) -> dict:
        return counts.setdefault(name or "-", {"drawings": 0, "added": 0, "revised": 0,
                                               "removed": 0, "released": 0, "on_hold": 0})

    for category, total in chain.totals_by_category.get(label, {}).items():
        bucket(category)["drawings"] = total

    if step is None:
        for category, entry in counts.items():
            entry["added"] = entry["drawings"]      # a first issue is all new
        return counts

    for field_name, values in (("added", step.added), ("removed", step.removed),
                               ("released", step.released)):
        for ident in values:
            bucket(split_member_id(ident)[0])[field_name] += 1
    for ident, _old, _new in step.revised:
        bucket(split_member_id(ident)[0])["revised"] += 1
    for hold in step.on_hold:
        bucket(hold.category or split_member_id(hold.ident)[0])["on_hold"] += 1
    return counts


def _write_summary(ws: Worksheet, chain: PackageChain) -> None:
    """The tracking table: one row per category per issue, then the issue total.

    A single line per issue answered "how much moved" but never "how much of
    what", and an assembly and a single part are not interchangeable work.
    """
    columns = ["#", "Code", "Stage", "Round", "Issue Folder", "Date", "Category",
               "Drawings", "Added", "Revised", "Removed", "Released", "On Hold",
               "Status"]
    _title(ws, "PACKAGE TRACKER", len(columns))

    row = 3
    for label, value in (
        ("Project", chain.project),
        ("Approved baseline", chain.baseline_label),
        ("Issues tracked", len(chain.issues)),
        ("Still on hold", len(chain.outstanding)),
        ("Generated", datetime.now().strftime("%d-%b-%Y %H:%M")),
    ):
        if value == "" or value is None:
            continue
        ws.cell(row=row, column=1, value=f"{label}:").font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value).font = _VALUE_FONT
        row += 1

    row += 1
    header_row = row
    _headers(ws, columns, header_row)
    row += 1

    # steps[i] describes the move into issues[i + 1]; the first issue has none.
    steps = [None] + list(chain.steps)
    for idx, entry in enumerate(chain.issues, start=1):
        step = steps[idx - 1] if idx - 1 < len(steps) else None
        is_iff = entry.stage == STAGE_IFF
        per_category = _category_counts(step, chain, entry.label)
        verdict = step.verdict if step else "first issue"

        lines = [(name, c, False) for name, c in per_category.items()]
        if len(lines) != 1:
            total = {k: sum(c[k] for c in per_category.values())
                     for k in ("drawings", "added", "revised", "removed",
                               "released", "on_hold")}
            lines.append(("All categories", total, True))

        for name, counts, is_total in lines:
            values = [
                idx, entry.code, entry.stage, entry.round_no, entry.label,
                entry.date_text, name,
                counts["drawings"], counts["added"], counts["revised"],
                counts["removed"],
                counts["released"] if is_iff else "",
                counts["on_hold"] if is_iff else "",
                verdict if is_total or len(lines) == 1 else "",
            ]
            for c_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c_idx, value=value)
                cell.border = _BORDER
                cell.alignment = _LEFT if c_idx in (5, 14) else _CENTER
                cell.fill = _IFF_FILL if is_iff else _IFA_FILL
                if is_total:
                    cell.font = _TOTAL_FONT
            row += 1

    row += 2
    row = _write_by_sequence(ws, chain, row)

    row += 1
    note = ws.cell(
        row=row, column=1,
        value="Members absent from an IFF release are on hold, not removed - "
              "fabrication ships the approved scope in slices. See the On Hold "
              "sheet for the reason recorded against each one.",
    )
    note.font = Font(italic=True, size=9, color="808080")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(columns))

    for idx, width in enumerate(
            [5, 9, 8, 8, 52, 13, 14, 10, 9, 9, 10, 10, 9, 34], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    # Below the header, wherever it landed - the label block above it is
    # shorter when a chain has no approved baseline yet. Set by coordinate:
    # a merged cell cannot be a freeze point.
    ws.freeze_panes = f"A{header_row + 1}"


def _write_by_sequence(ws: Worksheet, chain: PackageChain, row: int) -> int:
    """Where the package stands, one line per sequence. Returns the next free row.

    The issue table says what moved in each delivery. This says what is left:
    for every band of work, how much has been released and how much is still
    approved but unshipped. That is the question asked in a progress meeting,
    and answering it per sequence is what makes it actionable - a sequence is a
    slice the shop can finish.
    """
    tally: "OrderedDict[tuple, dict]" = OrderedDict()
    for key, info in chain.member_info.items():
        band = (info.get("band_kind", ""), info.get("band", ""))
        group = (info.get("zone", ""), band, info.get("category", ""))
        entry = tally.setdefault(group, {"members": 0, "released": 0, "held": 0,
                                         "dropped": 0})
        entry["members"] += 1
        if key in chain.released_keys:
            entry["released"] += 1
        elif key in chain.dropped_at:
            entry["dropped"] += 1
        hold = chain.holds.get(key)
        if hold is not None and not hold.released_in:
            entry["held"] += 1
    if not tally:
        return row

    heading = ws.cell(row=row, column=1, value="WHERE THE PACKAGE STANDS, BY SEQUENCE")
    heading.font = _LABEL_FONT
    row += 1

    # No "remaining" column: it is Members minus Released, which before the
    # first release equals Members and reads as a fourth bucket beside the
    # other three without being one.
    columns = ["Zone", "Sequence", "Category", "Members", "Released",
               "On Hold", "Dropped"]
    _headers(ws, columns, row)
    row += 1

    def order(group: tuple) -> tuple:
        zone, band, category = group
        zone_rank = (0, int(zone)) if str(zone).isdigit() else (1, 0)
        return (zone_rank, sequencing.band_sort_key(band), category)

    for group in sorted(tally, key=order):
        zone, band, category = group
        counts = tally[group]
        values = [zone or "-", sequencing.band_label(band) if band[1] else "-",
                  category or "-", counts["members"], counts["released"],
                  counts["held"], counts["dropped"]]
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c_idx, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if c_idx == 3 else _CENTER
            if counts["held"] and c_idx == 6:
                cell.fill = _HOLD_FILL
                cell.font = _HOLD_FONT
            elif counts["dropped"] and c_idx == 7:
                cell.fill = _DROP_FILL
                cell.font = _DROP_FONT
            elif counts["released"] == counts["members"]:
                cell.fill = _IFF_FILL          # this band is fully out
        row += 1

    totals = [
        "TOTAL", "", "",
        sum(c["members"] for c in tally.values()),
        sum(c["released"] for c in tally.values()),
        sum(c["held"] for c in tally.values()),
        sum(c["dropped"] for c in tally.values()),
    ]
    for c_idx, value in enumerate(totals, start=1):
        cell = ws.cell(row=row, column=c_idx, value=value)
        cell.border = _BORDER
        cell.font = _TOTAL_FONT
        cell.alignment = _LEFT if c_idx == 3 else _CENTER
    return row + 1


_LEGEND = [
    ("A", "Approval", "ifa"),
    ("B, C ..", "Re-Approval", "ifa"),
    ("H", "On Hold", "hold"),
    ("D", "Dropped At Re-Approval", "drop"),
    ("0", "Released For Fabrication", "iff"),
    ("1,2..", "Revised As Noted", "iff"),
    ("B / 1", "Revision Out Of Order", "rev"),
]

_LEGEND_FILLS = {"ifa": _IFA_FILL, "iff": _IFF_FILL, "hold": _HOLD_FILL,
                 "drop": _DROP_FILL, "rev": _REV_ERR_FILL}

# The one legend entry that needs more than two words: it marks a revision that
# is wrong rather than a state the member is in, so it says what was expected.
_REV_ERR_NOTE = ("A red revision skipped a rung: approval starts at A and steps "
                 "one letter per re-approval, fabrication starts at 0 and steps "
                 "one number per revised-as-noted. A red B means the A never "
                 "came - the issue that carried it is missing from this chain.")


def _write_legend(ws: Worksheet, column: int, row: int) -> None:
    """What the letters in the grid mean, beside the grid that uses them.

    The sheet is read by people who did not generate it, and a bare "B" or "0"
    in a coloured cell is not self-explanatory.
    """
    head = ws.cell(row=row, column=column, value="LEGEND")
    head.font = _LABEL_FONT
    row += 1
    for code, meaning, kind in _LEGEND:
        key = ws.cell(row=row, column=column, value=code)
        key.fill = _LEGEND_FILLS[kind]
        key.border = _BORDER
        key.alignment = _CENTER
        key.font = {"hold": _HOLD_FONT, "drop": _DROP_FONT,
                    "rev": _REV_ERR_FONT}.get(kind, _VALUE_FONT)
        text = ws.cell(row=row, column=column + 1, value=meaning)
        text.font = _VALUE_FONT
        text.alignment = _LEFT
        row += 1
    row += 1
    note = ws.cell(row=row, column=column, value=_REV_ERR_NOTE)
    note.font = Font(italic=True, size=9, color="808080")
    note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    # Six rows deep: a merged cell does not auto-fit, and the row heights are
    # shared with the member grid to the left, so the block is sized instead.
    ws.merge_cells(start_row=row, start_column=column,
                   end_row=row + 5, end_column=column + 1)
    ws.column_dimensions[get_column_letter(column)].width = 9
    ws.column_dimensions[get_column_letter(column + 1)].width = 26


def _write_history(ws: Worksheet, chain: PackageChain, category: str = "",
                   members: list[str] | None = None) -> None:
    """Member down the side, issue across the top, revision in the cell.

    This is the rev-by-rev view: one glance shows a member's whole life, when it
    entered the package, every revision it went through, and which release it
    shipped in.

    One sheet per drawing category, because an assembly and a single part are
    different deliverables that happen to share a mark - reading them in one
    list invites treating a cut part as a fabricated assembly.
    """
    ordered = members if members is not None else list(chain.history)
    held = {h.member_key for h in chain.outstanding}

    def was_in(entry) -> bool:
        """Does this issue have anything to say about this category?

        Single parts are detailed for fabrication, so an approval round carries
        none of them, and a column of blanks across every part reads as a
        package that dropped them all. An issue still earns its column when it
        is where a member was dropped, or a release that left one behind - the
        D and the H are what those columns are carrying.
        """
        if any(chain.history.get(k, {}).get(entry.label) for k in ordered):
            return True
        if any(chain.dropped_at.get(k) == entry.label for k in ordered):
            return True
        return entry.stage == STAGE_IFF and any(k in held for k in ordered)

    issues = [e for e in chain.issues if was_in(e)]
    labels = [e.label for e in issues]
    columns = ["Member Name", "Zone"] + [e.code for e in issues]
    heading = "MEMBER HISTORY - REVISION BY ISSUE"
    if category:
        heading += f"   ({category})"
    _title(ws, heading, max(len(columns), 3))

    _headers(ws, columns, 3)

    row = 4
    # Same banding as the register: sequence for assemblies, type for single
    # parts. A history sheet of 900 members read as one list hides which slice
    # of work each row belongs to.
    grouped: "OrderedDict[tuple[str, str], list[str]]" = OrderedDict()
    for key in ordered:
        info = chain.member_info.get(key, {})
        grouped.setdefault((info.get("band_kind", ""), info.get("band", "")),
                           []).append(key)
    banded = sorted(grouped, key=sequencing.band_sort_key)

    for band in banded:
        if band[1] and len(grouped) > 1:
            ws.merge_cells(start_row=row, start_column=1,
                           end_row=row, end_column=len(columns))
            head = ws.cell(row=row, column=1,
                           value=f"{sequencing.band_label(band)}"
                                 f"   -   {len(grouped[band])} member(s)")
            head.font = _BAND_FONT
            head.fill = _BAND_FILL
            head.alignment = _LEFT
            head.border = _BORDER
            row += 1
        row = _history_rows(ws, chain, grouped[band], issues, held, row)

    _write_legend(ws, len(columns) + 2, 3)

    ws.freeze_panes = "C4"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 8
    for idx in range(3, len(columns) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 11


def _history_rows(ws: Worksheet, chain: PackageChain, members: list[str],
                  issues: list, held: set, row: int) -> int:
    """The member rows of one band. Returns the next free row."""
    labels = [e.label for e in issues]
    for member in members:
        per_issue = chain.history.get(member, {})
        faults = chain.rev_errors.get(member, {})
        info = chain.member_info.get(member, {})
        name = info.get("name") or split_member_id(member)[1]
        zone = info.get("zone", "")
        cells: list[object] = [name, zone]
        for label in labels:
            cells.append(per_issue.get(label, ""))
        for c_idx, value in enumerate(cells, start=1):
            cell = ws.cell(row=row, column=c_idx, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if c_idx == 1 else _CENTER
            if c_idx > 2:
                entry = issues[c_idx - 3]
                if value and faults.get(entry.label):
                    # The revision itself is still printed - the reader needs to
                    # see that it is a B to see that the A never came. What red
                    # means is in the legend beside the grid, not in a comment
                    # that has to be hovered to be found.
                    cell.fill = _REV_ERR_FILL
                    cell.font = _REV_ERR_FONT
                elif value:
                    cell.fill = _IFF_FILL if entry.stage == STAGE_IFF else _IFA_FILL
                elif entry.stage == STAGE_IFF and member in held:
                    # Say it, do not just colour it: a printed tracker and a
                    # colour-blind reader both lose a fill, and "on hold" is the
                    # one state the shop floor acts on.
                    cell.value = "H"
                    cell.fill = _HOLD_FILL
                    cell.font = _HOLD_FONT
                elif chain.dropped_at.get(member) == entry.label:
                    # Delivered and dropped were both blank white, so a member
                    # withdrawn at re-approval read exactly like one that had
                    # shipped. This is the issue that withdrew it.
                    cell.value = "D"
                    cell.fill = _DROP_FILL
                    cell.font = _DROP_FONT
        row += 1

    return row


def _write_holds(ws: Worksheet, chain: PackageChain) -> None:
    """Approved members that have not shipped, and why."""
    columns = ["Member Name", "Category", "Zone", "Sequence", "Approved Rev",
               "On Hold Since", "Released In", "Reason"]
    _title(ws, "ON HOLD - APPROVED BUT NOT YET RELEASED", len(columns))
    _headers(ws, columns, 3)

    # Ordered by band, so the sheet reads as the slices of work it is - but as
    # a sortable column rather than heading rows, because this sheet is looked
    # things up in ("which of these still has no reason?") and heading rows
    # would cost it the filter.
    row = 4
    ordered = sorted(chain.holds.values(),
                     key=lambda h: (sequencing.band_sort_key(chain.band_of(h.ident)),
                                    h.member_name))
    for hold in ordered:
        shipped = bool(hold.released_in)
        band = chain.band_of(hold.ident)
        values = [hold.member_name, hold.category or "-", hold.zone,
                  sequencing.band_label(band) if band[1] else "-",
                  hold.revision, hold.held_since,
                  hold.released_in or "", hold.reason or ""]
        for c_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c_idx, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if c_idx in (1, 8) else _CENTER
            if shipped:
                cell.fill = _OK_FILL
            elif c_idx == 8 and not hold.reason:
                cell.fill = _REVIEW_FILL   # amber: nobody has said why yet
            else:
                cell.fill = _HOLD_FILL
        row += 1

    ws.freeze_panes = "A4"
    if row > 4:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{row - 1}"
    for idx, width in enumerate([24, 12, 8, 14, 14, 30, 30, 52], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_changes(ws: Worksheet, chain: PackageChain) -> None:
    """Every change, issue by issue - the comparison log, flattened."""
    columns = ["Issue", "Code", "Stage", "Change", "Category", "Sequence",
               "Member Name", "From Rev", "To Rev"]
    _title(ws, "CHANGE LOG", len(columns))
    _headers(ws, columns, 3)

    row = 4
    for step in chain.steps:
        groups: list[tuple[str, list, PatternFill]] = [
            ("Added", [(m, "", "") for m in step.added], _NEW_FILL),
            ("Revised", list(step.revised), _REVIEW_FILL),
            ("Removed", [(m, "", "") for m in step.removed], _HOLD_FILL),
            ("Released", [(m, "", "") for m in step.released], _IFF_FILL),
            ("On Hold", [(h.ident, h.revision, "") for h in step.on_hold], _HOLD_FILL),
        ]
        for change, items, fill in groups:
            # Within a change, in sequence order: the log is read looking for
            # what happened to one slice of work.
            items = sorted(items, key=lambda i: (
                sequencing.band_sort_key(chain.band_of(i[0])), i[0]))
            for item in items:
                ident, from_rev, to_rev = item
                # The lists carry identities, so the same mark issued as both an
                # assembly and a part reads as the two separate rows it is.
                category, name = split_member_id(ident)
                band = chain.band_of(ident)
                values = [step.new_label, step.code, step.stage, change,
                          category, sequencing.band_label(band) if band[1] else "-",
                          name, from_rev, to_rev]
                for c_idx, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=c_idx, value=value)
                    cell.border = _BORDER
                    cell.alignment = _LEFT if c_idx in (1, 7) else _CENTER
                    cell.fill = fill
                row += 1

    ws.freeze_panes = "A4"
    if row > 4:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{row - 1}"
    for idx, width in enumerate([46, 10, 8, 12, 14, 14, 24, 11, 11], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _by_category(chain: PackageChain) -> list[tuple[str, list[str]]]:
    """Members grouped by drawing category, in the order they were first seen."""
    grouped: "OrderedDict[str, list[str]]" = OrderedDict()
    for ident in chain.history:
        grouped.setdefault(split_member_id(ident)[0], []).append(ident)
    if not grouped:
        return [("", [])]
    return list(grouped.items())


def write_tracker(chain: PackageChain, output_path: str | Path) -> Path:
    """Write the cumulative tracking workbook for one package."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    used: set[str] = set()
    _write_summary(wb.create_sheet(safe_sheet_name("Tracker", used)), chain)

    # One history sheet per drawing category, in the same workbook. A package
    # that issues assembly and single-part drawings is one package with one
    # chain - splitting it across two tracker files would break the chaining
    # that the tracker exists for.
    for category, members in _by_category(chain):
        title = f"History - {category}" if category else "Member History"
        _write_history(wb.create_sheet(safe_sheet_name(title, used)), chain,
                       category, members)

    _write_changes(wb.create_sheet(safe_sheet_name("Change Log", used)), chain)
    if chain.holds:
        _write_holds(wb.create_sheet(safe_sheet_name("On Hold", used)), chain)

    _save_workbook(wb, out)
    return out


def suggest_tracker_name(chain: PackageChain) -> str:
    """Filename for the tracker, named after the package it follows."""
    stem = (chain.project or "Package").strip()
    safe = "".join(ch for ch in stem if ch not in '<>:"/\\|?*')
    safe = safe[:110].strip(" .")
    return f"{safe or 'Package'} - Tracker.xlsx"
