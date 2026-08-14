"""Tests for the IFA/IFF package chain.

The rules under test are the ones that differ from a pairwise register diff:
absent members are on hold rather than removed once fabrication starts, and
releases accumulate against the approved baseline instead of against each other.
"""

from __future__ import annotations

from collections import OrderedDict
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import make_drawing

from fabdoc.register import build_register
from fabdoc.tracking import (STAGE_IFA, STAGE_IFF, ChainState, IssueEntry,
                             apply_reasons, build_chain, load_state,
                             project_name_from, save_state, snapshot_register,
                             state_path_for)
from fabdoc.tracking_out import write_tracker


def key(mark: str, category: str = "Assembly") -> str:
    """The id the chain tracks a mark under, now that category is part of it."""
    return f"{category}::{mark}"


def issue(label: str, stage: str, round_no: str, marks: dict[str, str]) -> IssueEntry:
    members: "OrderedDict[str, dict[str, str]]" = OrderedDict()
    for mark, rev in marks.items():
        members[mark] = {"rev": rev, "zone": mark[2] if len(mark) > 2 else "", "category": "Assembly"}
    return IssueEntry(label=label, stage=stage, round_no=round_no, members=members)


def approved_then_released(released: list[dict[str, str]]) -> ChainState:
    """An approved scope of five members, then the given IFF releases."""
    state = ChainState(project="Stairs")
    state.add_issue(issue("IFA-1", STAGE_IFA, "1",
                          {"A1": "A", "A2": "A", "A3": "A", "A4": "A", "A5": "A"}))
    for idx, marks in enumerate(released, start=1):
        state.add_issue(issue(f"IFF-{idx}", STAGE_IFF, str(idx), marks))
    return state


# ------------------------------------------------------------------ IFA rounds


def test_ifa_round_reports_added_revised_and_removed():
    state = ChainState(project="Stairs")
    state.add_issue(issue("IFA-15", STAGE_IFA, "15", {"A1": "A", "A2": "A", "A3": "A"}))
    state.add_issue(issue("IFA-25", STAGE_IFA, "25", {"A1": "A", "A2": "B", "A9": "A"}))
    step = build_chain(state).steps[-1]
    assert step.added == [key("A9")]
    assert step.removed == [key("A3")]
    assert [r[0] for r in step.revised] == [key("A2")]
    assert step.unchanged == [key("A1")]


def test_within_approval_a_dropped_drawing_is_still_removed():
    """The on-hold rule must not leak backwards into the IFA rounds."""
    state = ChainState()
    state.add_issue(issue("IFA-1", STAGE_IFA, "1", {"A1": "A", "A2": "A"}))
    state.add_issue(issue("IFA-2", STAGE_IFA, "2", {"A1": "B"}))
    step = build_chain(state).steps[-1]
    assert step.removed == [key("A2")]
    assert not step.on_hold


# ------------------------------------------------------- IFF partial releases


def test_unshipped_members_are_on_hold_not_removed():
    chain = build_chain(approved_then_released([{"A1": "A", "A2": "A"}]))
    step = chain.steps[-1]
    assert step.removed == []
    assert sorted(h.member_name for h in step.on_hold) == ["A3", "A4", "A5"]
    assert sorted(step.released) == [key("A1"), key("A2")]


def test_second_release_does_not_mark_the_first_as_removed():
    """The trap a pairwise diff falls into: IFF-2 holds the next slice, not the previous."""
    chain = build_chain(approved_then_released([{"A1": "A", "A2": "A"}, {"A3": "A"}]))
    step = chain.steps[-1]
    assert step.removed == []
    assert step.released == [key("A3")]
    assert sorted(h.member_name for h in step.on_hold) == ["A4", "A5"]


def test_outstanding_shrinks_with_every_release():
    counts = []
    for releases in ([{"A1": "A"}],
                     [{"A1": "A"}, {"A2": "A", "A3": "A"}],
                     [{"A1": "A"}, {"A2": "A", "A3": "A"}, {"A4": "A", "A5": "A"}]):
        counts.append(len(build_chain(approved_then_released(releases)).outstanding))
    assert counts == [4, 2, 0]


def test_a_held_member_records_where_it_finally_shipped():
    chain = build_chain(approved_then_released([{"A1": "A"}, {"A3": "A"}]))
    hold = chain.holds[key("A3")]
    assert hold.held_since == "IFF-1"
    assert hold.released_in == "IFF-2"
    assert hold not in chain.outstanding


def test_a_release_may_add_scope_the_baseline_never_had():
    chain = build_chain(approved_then_released([{"A1": "A", "NEW9": "A"}]))
    step = chain.steps[-1]
    assert step.added == [key("NEW9")]
    assert step.released == [key("A1")]


def test_fabrication_without_a_prior_approval_uses_itself_as_baseline():
    state = ChainState()
    state.add_issue(issue("IFF-1", STAGE_IFF, "1", {"A1": "A"}))
    chain = build_chain(state)
    assert chain.outstanding == []
    assert chain.steps[-1].released == [key("A1")]


# -------------------------------------------------------------- hold reasons


def test_outstanding_members_start_without_a_reason():
    chain = build_chain(approved_then_released([{"A1": "A"}]))
    assert len(chain.missing_reasons()) == 4


def test_a_batch_reason_can_be_overridden_per_member():
    state = approved_then_released([{"A1": "A"}])
    chain = build_chain(state)
    # Keyed by identity, not by the displayed mark: that is what a dialog
    # hands back, and what keeps an assembly and a part apart.
    reasons = {h.ident: "Client hold" for h in chain.outstanding}
    reasons[key("A5")] = "Material shortage"
    apply_reasons(state, reasons)
    chain = build_chain(state)
    assert chain.missing_reasons() == []
    assert chain.holds[key("A5")].reason == "Material shortage"
    assert chain.holds[key("A4")].reason == "Client hold"


def test_a_reason_survives_into_later_releases():
    state = approved_then_released([{"A1": "A"}])
    apply_reasons(state, {key("A5"): "Client hold"})
    state.add_issue(issue("IFF-2", STAGE_IFF, "2", {"A2": "A"}))
    chain = build_chain(state)
    assert chain.holds[key("A5")].reason == "Client hold"


def test_blank_reasons_are_ignored():
    state = approved_then_released([{"A1": "A"}])
    apply_reasons(state, {key("A5"): "   "})
    assert build_chain(state).holds[key("A5")].needs_reason


# ------------------------------------------------------------- state and names


def test_reprocessing_a_folder_updates_it_rather_than_chaining_a_duplicate():
    state = ChainState()
    state.add_issue(issue("IFA-1", STAGE_IFA, "1", {"A1": "A"}))
    state.add_issue(issue("IFA-1", STAGE_IFA, "1", {"A1": "A", "A2": "A"}))
    assert len(state.issues) == 1
    assert state.issues[0].total == 2


def test_chain_state_round_trips_through_json(tmp_path: Path):
    state = approved_then_released([{"A1": "A"}])
    apply_reasons(state, {key("A5"): "Client hold"})
    path = state_path_for(tmp_path / "Package Tracker.xlsx")
    save_state(state, path)
    assert path.name.endswith(".chain.json")

    reloaded = load_state(path)
    assert [e.label for e in reloaded.issues] == ["IFA-1", "IFF-1"]
    assert build_chain(reloaded).holds[key("A5")].reason == "Client hold"


def test_missing_state_file_yields_an_empty_chain(tmp_path: Path):
    assert load_state(tmp_path / "nope.chain.json").issues == []


@pytest.mark.parametrize("title", [
    "Stairs at Zone 1 and Zone 2 for Approval",
    "Stairs at Zone 1 and Zone 2 for Re Approval",
    "Stairs at Zone 1 and Zone 2 for Re-Approval",
    "Stairs at Zone 1 and Zone 2 for Fabrication",
    # A stage-coded folder prefix survives folder_meta, which only lifts a
    # numeric "25." into issue_no. Left in, the fabrication issues chained into
    # their own tracker and the on-hold rule silently switched itself off.
    "IFF-0    Stairs at Zone 1 and Zone 2 for Fabrication",
    "IFF-3 Stairs at Zone 1 and Zone 2 for Fabrication",
    "IFA-2 Stairs at Zone 1 and Zone 2 for Re Approval",
    "IFF 1 - Stairs at Zone 1 and Zone 2",
])
def test_every_issue_purpose_maps_to_one_project(title: str):
    assert project_name_from(title) == "Stairs at Zone 1 and Zone 2"


def test_a_title_without_a_purpose_clause_is_left_alone():
    assert project_name_from("Skyline Tower") == "Skyline Tower"


# ------------------------------------------------------------------- workbook


def test_tracker_workbook_has_the_expected_sheets(tmp_path: Path):
    state = approved_then_released([{"A1": "A"}])
    apply_reasons(state, {key("A5"): "Client hold"})
    out = write_tracker(build_chain(state), tmp_path / "tracker.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Tracker", "History - Assembly", "Change Log", "On Hold"]
    wb.close()


def test_member_history_shows_a_revision_per_issue(tmp_path: Path):
    state = ChainState()
    state.add_issue(issue("IFA-1", STAGE_IFA, "1", {"A1": "A"}))
    state.add_issue(issue("IFA-2", STAGE_IFA, "2", {"A1": "B"}))
    out = write_tracker(build_chain(state), tmp_path / "tracker.xlsx")
    wb = load_workbook(out)
    ws = wb["History - Assembly"]
    # The legend sits to the right of the table, so read the table's columns.
    assert [c.value for c in ws[3][:4]] == ["Member Name", "Zone", "IFA-1", "IFA-2"]
    row = [c.value for c in ws[4][:4]]
    assert row[0] == "A1" and row[2:] == ["A", "B"]      # zone is unset here
    wb.close()


def test_on_hold_sheet_carries_the_reason(tmp_path: Path):
    state = approved_then_released([{"A1": "A"}])
    apply_reasons(state, {key(h): "Client hold" for h in ("A2", "A3", "A4", "A5")})
    out = write_tracker(build_chain(state), tmp_path / "tracker.xlsx")
    wb = load_workbook(out)
    ws = wb["On Hold"]
    header = [c.value for c in ws[3]]
    reason = header.index("Reason")
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert {r[0] for r in rows if r[0]} == {"A2", "A3", "A4", "A5"}
    assert all(r[reason] == "Client hold" for r in rows if r[0])
    wb.close()


# ------------------------------------------------------- against real drawings


def test_snapshot_keeps_one_row_per_member(tmp_path: Path):
    root = tmp_path / "15. Stairs at Zone 1 for Approval - 2026-06-26"
    for mark in ("17172C1", "17172C2", "17271X1"):
        make_drawing(root / "Assembly" / f"{mark}  - Rev A.pdf", mark, revision="A")
    snap = snapshot_register(build_register(root))
    # Keyed by category and mark, because a part drawing may carry the same one.
    assert set(snap) == {key("17172C1"), key("17172C2"), key("17271X1")}
    assert snap[key("17271X1")]["rev"] == "A"
    assert snap[key("17271X1")]["zone"] == "2"
    assert snap[key("17271X1")]["name"] == "17271X1"


# ------------------------------------------------- categories in one tracker


def part(label: str, stage: str, round_no: str,
         marks: dict[str, str], category: str) -> IssueEntry:
    members: "OrderedDict[str, dict[str, str]]" = OrderedDict()
    for mark, rev in marks.items():
        members[mark] = {"rev": rev, "zone": "1", "category": category}
    return IssueEntry(label=label, stage=stage, round_no=round_no, members=members)


def test_a_part_and_an_assembly_sharing_a_mark_are_two_items():
    """Keyed by mark alone, whichever was read second overwrote the first.

    An assembly drawing is fabricated and a single-part drawing is cut; they
    are different deliverables that routinely carry the same mark.
    """
    members: "OrderedDict[str, dict[str, str]]" = OrderedDict()
    members["17172C172"] = {"rev": "A", "category": "Assembly"}
    members["Part::17172C172"] = {"rev": "0", "category": "Part"}
    entry = IssueEntry(label="IFA-1", stage=STAGE_IFA, round_no="1", members=members)

    assert entry.total == 2
    assert set(entry.members) == {"Assembly::17172C172", "Part::17172C172"}
    assert entry.members["Part::17172C172"]["name"] == "17172C172"


def test_a_part_is_not_released_by_shipping_the_assembly():
    state = ChainState(project="Zone 1")
    state.add_issue(part("IFA-1", STAGE_IFA, "1", {"C172": "A"}, "Assembly"))
    state.issues[-1].members["Part::C172"] = {
        "name": "C172", "rev": "0", "zone": "1", "category": "Part"}
    state.issues[-1].__post_init__()
    state.add_issue(part("IFF-2", STAGE_IFF, "2", {"C172": "A"}, "Assembly"))

    chain = build_chain(state)
    step = chain.steps[-1]
    assert step.released == ["Assembly::C172"]
    # The part of the same mark has not shipped and is still owed.
    assert [h.ident for h in step.on_hold] == ["Part::C172"]
    assert [h.category for h in chain.outstanding] == ["Part"]


def test_every_category_shares_one_tracker_file(tmp_path: Path):
    """Parts get their own sheet, not their own workbook.

    A package that issues assemblies and single parts is one package with one
    chain; a second file would break the chaining the tracker exists for.
    """
    state = ChainState(project="Zone 1")
    entry = part("IFA-1", STAGE_IFA, "1", {"C1": "A", "C2": "A"}, "Assembly")
    entry.members["Part::P9"] = {"name": "P9", "rev": "0", "zone": "1",
                                 "category": "Part"}
    entry.__post_init__()
    state.add_issue(entry)

    out = write_tracker(build_chain(state), tmp_path / "tracker.xlsx")
    assert list(tmp_path.glob("*.xlsx")) == [out], "a second workbook was written"

    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["Tracker", "History - Assembly", "History - Part",
                                 "Change Log"]
        assert _member_names(wb["History - Part"]) == ["P9"]
        assert sorted(_member_names(wb["History - Assembly"])) == ["C1", "C2"]
    finally:
        wb.close()


def test_reasons_can_be_keyed_by_identity_or_by_mark():
    """A dialog hands back whatever it used as a row id."""
    state = ChainState(project="Zone 1")
    state.add_issue(part("IFA-1", STAGE_IFA, "1", {"A1": "A", "A2": "A"}, "Assembly"))
    state.add_issue(part("IFF-2", STAGE_IFF, "2", {"A1": "A"}, "Assembly"))

    apply_reasons(state, {"Assembly::A2": "Client hold"})
    assert build_chain(state).holds["Assembly::A2"].reason == "Client hold"


# --------------------------------------------------------------- mark casing


def _member_names(ws) -> list[str]:
    """Member column only.

    Skips the legend rows to the right and the band headings between clusters,
    which both live in the same rows as member data.
    """
    return [r[0] for r in ws.iter_rows(min_row=4, values_only=True)
            if r and r[0] and "member(s)" not in str(r[0])]


def _members(names, category="Part", rev="0"):
    from collections import OrderedDict
    return OrderedDict(
        (n, {"name": n, "rev": rev, "zone": "", "category": category}) for n in names
    )


def test_a_case_change_between_issues_is_the_same_member():
    """Marks are now recorded as drawn, and existing trackers hold them upper.

    An issue added after that change must chain onto the ones before it, not
    read as the old member removed and a new one added.
    """
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["17CH104", "17HSP1"])))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["17ch104"])))
    chain = build_chain(state)
    step = chain.steps[-1]

    assert len(step.released) == 1
    assert not step.added, "a case change read as a different member"
    assert len(step.on_hold) == 1


def test_a_case_change_does_not_duplicate_the_history_row():
    """Member History is keyed by comparison key, so one member is one row."""
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["17CH104"])))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["17ch104"])))
    chain = build_chain(state)

    assert len(chain.history) == 1
    # The most recent spelling is what the sheet shows.
    assert list(chain.member_info.values())[0]["name"] == "17ch104"


def test_a_hold_reason_survives_a_case_change():
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["17ch104", "17hsp1"])))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["17ch104"])))
    chain = build_chain(state)
    hold = chain.outstanding[0]
    assert hold.member_name == "17hsp1"

    apply_reasons(state, {hold.ident: "awaiting material"})
    assert build_chain(state).outstanding[0].reason == "awaiting material"


def test_the_history_sheet_lists_each_member_once(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["17CH104", "17HSP1"])))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["17ch104"])))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")

    wb = load_workbook(out)
    try:
        sheet = next(n for n in wb.sheetnames if "History" in n)
        names = _member_names(wb[sheet])
    finally:
        wb.close()
    assert sorted(names) == ["17HSP1", "17ch104"]


# ----------------------------------------------- legend, H, untracked kinds


def test_an_on_hold_cell_says_H_as_well_as_being_orange(tmp_path: Path):
    """A printed tracker loses the fill, and so does a colour-blind reader.

    On hold is the one state the shop floor acts on, so it is written down.
    """
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["A1", "A2"], category="Assembly")))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["A1"], category="Assembly")))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")

    wb = load_workbook(out)
    try:
        ws = wb["History - Assembly"]
        rows = {r[0]: r[1:4] for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}
        assert rows["A1"][2] == "0", "released member keeps its revision"
        assert rows["A2"][2] == "H", "held member is marked H"
    finally:
        wb.close()


def test_the_history_sheet_carries_a_legend(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["A1"], category="Assembly")))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")

    wb = load_workbook(out)
    try:
        ws = wb["History - Assembly"]
        seen = {c.value for row in ws.iter_rows() for c in row if c.value}
        assert "LEGEND" in seen
        for meaning in ("Approval", "Re-Approval", "On Hold",
                        "Released For Fabrication", "Revised As Noted"):
            assert meaning in seen, meaning
    finally:
        wb.close()


def test_erection_drawings_are_not_tracked(tmp_path: Path):
    """An erection drawing shows where an assembly goes; it is not fabricated.

    It has no approved scope to release and nothing to hold, so counting it put
    site drawings into the released and on-hold numbers the shop reads.
    """
    members = _members(["A1"], category="Assembly")
    members.update(_members(["E1", "E2"], category="Erection"))
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=members))
    chain = build_chain(state)

    assert list(chain.history) == ["Assembly::A1"]
    assert chain.totals["approval"] == 1, "the Drawings count follows what is tracked"

    out = write_tracker(chain, tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        assert "History - Erection" not in wb.sheetnames
    finally:
        wb.close()


def test_dropping_erection_does_not_report_it_as_removed():
    """Existing chains hold erection drawings; excluding them is not a deletion."""
    first = _members(["A1"], category="Assembly")
    first.update(_members(["E1", "E2"], category="Erection"))
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1", members=first))
    state.add_issue(IssueEntry(label="two", stage=STAGE_IFA, round_no="2",
                               members=_members(["A1"], category="Assembly")))
    step = build_chain(state).steps[-1]
    assert step.removed == [], "erection drawings read as removed"
    assert step.verdict == "no change"


def test_the_untracked_list_is_editable():
    from fabdoc.config import AppSettings

    cfg = AppSettings()
    cfg.untracked_categories = ["Part"]
    members = _members(["A1"], category="Assembly")
    members.update(_members(["P1"], category="Part"))
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1",
                               members=members))
    assert list(build_chain(state, cfg).history) == ["Assembly::A1"]


def _tracker_rows(ws) -> list[tuple]:
    """Data rows of the Tracker sheet: (#, code, category, drawings, added,
    revised, removed, released, on_hold)."""
    rows = []
    seen_header = False
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "#":
            seen_header = True
            continue
        if not seen_header or not r or r[0] is None:
            continue
        # The by-sequence table and the footnote follow the issue table.
        if isinstance(r[0], str) and (r[0].startswith("Members absent")
                                      or r[0].startswith("WHERE THE PACKAGE")
                                      or r[0] == "Zone"):
            break
        rows.append((r[0], r[1], r[6], r[7], r[8], r[9], r[10], r[11], r[12]))
    return rows


def test_the_tracking_table_splits_each_issue_by_category(tmp_path: Path):
    """An assembly and a single part are not interchangeable work.

    "324 drawings released" says nothing useful until it says how much of it
    was assemblies.
    """
    approved = _members(["A1", "A2"], category="Assembly")
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=approved))

    release = _members(["A1"], category="Assembly")
    release.update(_members(["P1", "P2", "P3"], category="Part"))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=release))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _tracker_rows(wb["Tracker"])
    finally:
        wb.close()

    by_name = {(r[0], r[2]): r for r in rows}
    assert by_name[(2, "Assembly")][3] == 1        # drawings
    assert by_name[(2, "Assembly")][7] == 1        # released
    assert by_name[(2, "Assembly")][8] == 1        # on hold (A2)
    assert by_name[(2, "Part")][3] == 3
    assert by_name[(2, "Part")][4] == 3            # all new to the package

    total = by_name[(2, "All categories")]
    assert total[3] == 4, "the total must equal the category rows"
    assert total[4] == 3 and total[7] == 1 and total[8] == 1


def test_the_category_rows_sum_to_the_total_row(tmp_path: Path):
    approved = _members(["A1", "A2", "A3"], category="Assembly")
    approved.update(_members(["P1"], category="Part"))
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=approved))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_members(["A1", "P1"], category="Assembly")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _tracker_rows(wb["Tracker"])
    finally:
        wb.close()

    for issue_no in {r[0] for r in rows}:
        parts = [r for r in rows if r[0] == issue_no and r[2] != "All categories"]
        totals = [r for r in rows if r[0] == issue_no and r[2] == "All categories"]
        if not totals:
            continue
        for col in range(3, 9):
            # Released and On Hold are left blank on an approval issue, so a
            # blank total is a zero, not a mismatch.
            summed = sum(r[col] for r in parts if isinstance(r[col], int))
            assert (totals[0][col] or 0) == summed, f"column {col} of issue {issue_no}"


def test_a_single_category_issue_stays_one_row(tmp_path: Path):
    """No "All categories" line when there is only one category to total."""
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_members(["A1"], category="Assembly")))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _tracker_rows(wb["Tracker"])
    finally:
        wb.close()
    assert [r[2] for r in rows] == ["Assembly"]
    assert rows[0][3] == 1


def _seq_summary(ws) -> list[tuple]:
    """Rows are (zone, sequence, category, members, released, held, dropped)."""
    """The by-sequence table on the Tracker sheet."""
    rows, on = [], False
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith("WHERE THE PACKAGE"):
            on = True
            continue
        if not on or not r or r[0] is None or r[0] == "Zone":
            continue
        if isinstance(r[0], str) and r[0].startswith("Members absent"):
            break
        rows.append(tuple(r[:7]))
    return rows


def _banded(names: list[str], category="Assembly", rev="A"):
    from collections import OrderedDict
    return OrderedDict(
        (n, {"name": n, "rev": rev, "zone": n[2] if len(n) > 2 else "",
             "category": category}) for n in names
    )


def test_the_history_sheet_bands_members_by_sequence(tmp_path: Path):
    """The same banding as the register: a flat list hides the slices of work."""
    state = ChainState(project="P")
    state.add_issue(IssueEntry(
        label="approval", stage=STAGE_IFA, round_no="1",
        members=_banded(["17172C1", "17172C2", "17173R1"])))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")

    wb = load_workbook(out)
    try:
        ws = wb["History - Assembly"]
        heads = [str(r[0]) for r in ws.iter_rows(min_row=4, values_only=True)
                 if r and r[0] and "member(s)" in str(r[0])]
    finally:
        wb.close()
    assert heads == ["SEQ 172   -   2 member(s)", "SEQ 173   -   1 member(s)"]


def test_single_parts_band_by_type_in_the_tracker_too(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(
        label="approval", stage=STAGE_IFA, round_no="1",
        members=_banded(["17ch1", "17ch2", "17hsp9"], category="Part")))
    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")

    wb = load_workbook(out)
    try:
        heads = [str(r[0]) for r in wb["History - Part"].iter_rows(
            min_row=4, values_only=True) if r and r[0] and "member(s)" in str(r[0])]
    finally:
        wb.close()
    assert heads == ["TYPE CH   -   2 member(s)", "TYPE HSP   -   1 member(s)"]


def test_the_main_sheet_reports_where_each_sequence_stands(tmp_path: Path):
    """The issue table says what moved; this says what is left, per sequence."""
    state = ChainState(project="P")
    state.add_issue(IssueEntry(
        label="approval", stage=STAGE_IFA, round_no="1",
        members=_banded(["17172C1", "17172C2", "17173R1"])))
    state.add_issue(IssueEntry(
        label="release", stage=STAGE_IFF, round_no="1",
        members=_banded(["17172C1"], rev="0")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _seq_summary(wb["Tracker"])
    finally:
        wb.close()

    by_band = {r[1]: r for r in rows if r[0] != "TOTAL"}
    # members, released, on hold, dropped
    assert by_band["SEQ 172"][3:] == (2, 1, 1, 0)
    assert by_band["SEQ 173"][3:] == (1, 0, 1, 0)

    total = next(r for r in rows if r[0] == "TOTAL")
    assert total[3:] == (3, 1, 2, 0)


def test_the_by_sequence_totals_match_the_rows(tmp_path: Path):
    state = ChainState(project="P")
    members = _banded(["17172C1", "17173R1"])
    members.update(_banded(["17ch1"], category="Part"))
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=members))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_banded(["17172C1"], rev="0")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _seq_summary(wb["Tracker"])
    finally:
        wb.close()

    body = [r for r in rows if r[0] != "TOTAL"]
    total = next(r for r in rows if r[0] == "TOTAL")
    for col in range(3, 7):
        assert total[col] == sum(r[col] for r in body), f"column {col}"


def _sheet(ws) -> tuple[list[str], list[tuple]]:
    header = [c.value for c in ws[3]]
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r and r[0]]
    return header, rows


def test_the_on_hold_sheet_carries_and_orders_by_sequence(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(
        label="approval", stage=STAGE_IFA, round_no="1",
        members=_banded(["17173R1", "17172C2", "17172C1"])))
    state.add_issue(IssueEntry(
        label="release", stage=STAGE_IFF, round_no="1",
        members=_banded(["17172C1"], rev="0")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        header, rows = _sheet(wb["On Hold"])
    finally:
        wb.close()

    seq = header.index("Sequence")
    # 17172C1 shipped in the same release, so it never went on hold at all.
    assert [r[0] for r in rows] == ["17172C2", "17173R1"]
    assert [r[seq] for r in rows] == ["SEQ 172", "SEQ 173"]


def test_the_change_log_carries_and_orders_by_sequence(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1"])))
    state.add_issue(IssueEntry(
        label="two", stage=STAGE_IFA, round_no="2",
        members=_banded(["17172C1", "17173R9", "17172C4"])))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        header, rows = _sheet(wb["Change Log"])
    finally:
        wb.close()

    seq, name = header.index("Sequence"), header.index("Member Name")
    added = [(r[seq], r[name]) for r in rows if r[3] == "Added"]
    assert added == [("SEQ 172", "17172C4"), ("SEQ 173", "17173R9")]


def test_both_flat_sheets_stay_filterable(tmp_path: Path):
    """Sequence is a column, not a heading row, so the filter survives.

    These two sheets are looked things up in - "which of these still has no
    reason?" - and heading rows would cost them that.
    """
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1", "17173R1"])))
    state.add_issue(IssueEntry(label="two", stage=STAGE_IFF, round_no="1",
                               members=_banded(["17172C1"], rev="0")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        for sheet in ("On Hold", "Change Log"):
            assert wb[sheet].auto_filter.ref, f"{sheet} lost its filter"
            assert "Sequence" in [c.value for c in wb[sheet][3]]
    finally:
        wb.close()


def test_a_member_dropped_at_re_approval_is_marked_D(tmp_path: Path):
    """Delivered and dropped were both a blank white cell.

    One is finished work, the other fell out of scope, and on a 900-member
    sheet the difference is the question "is this done, or did we lose it?".
    """
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1", "17172C2"])))
    state.add_issue(IssueEntry(label="re-approval", stage=STAGE_IFA, round_no="2",
                               members=_banded(["17172C1"], rev="B")))
    chain = build_chain(state)
    assert chain.dropped_at == {"Assembly::17172C2": "re-approval"}

    out = write_tracker(chain, tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        ws = wb["History - Assembly"]
        rows = {r[0]: r[2:4] for r in ws.iter_rows(min_row=4, values_only=True)
                if r and r[0] and "member(s)" not in str(r[0])}
        seen = {c.value for row in ws.iter_rows() for c in row if c.value}
    finally:
        wb.close()

    assert rows["17172C2"] == ("A", "D"), "the issue that dropped it is marked"
    assert rows["17172C1"] == ("A", "B")
    assert "Dropped At Re-Approval" in seen, "the legend explains D"


def test_a_member_that_comes_back_is_not_marked_dropped(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1", "17172C2"])))
    state.add_issue(IssueEntry(label="two", stage=STAGE_IFA, round_no="2",
                               members=_banded(["17172C1"], rev="B")))
    state.add_issue(IssueEntry(label="three", stage=STAGE_IFA, round_no="3",
                               members=_banded(["17172C1", "17172C2"], rev="C")))
    assert build_chain(state).dropped_at == {}


def test_absence_from_a_release_is_a_hold_not_a_drop():
    """Only an approval round drops a member; a release holds the balance."""
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="approval", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1", "17172C2"])))
    state.add_issue(IssueEntry(label="release", stage=STAGE_IFF, round_no="1",
                               members=_banded(["17172C1"], rev="0")))
    chain = build_chain(state)
    assert chain.dropped_at == {}
    assert [h.member_name for h in chain.outstanding] == ["17172C2"]


def test_released_plus_held_plus_dropped_accounts_for_every_member(tmp_path: Path):
    state = ChainState(project="P")
    state.add_issue(IssueEntry(label="one", stage=STAGE_IFA, round_no="1",
                               members=_banded(["17172C1", "17172C2", "17172C3"])))
    state.add_issue(IssueEntry(label="two", stage=STAGE_IFA, round_no="2",
                               members=_banded(["17172C1", "17172C2"], rev="B")))
    state.add_issue(IssueEntry(label="three", stage=STAGE_IFF, round_no="1",
                               members=_banded(["17172C1"], rev="0")))

    out = write_tracker(build_chain(state), tmp_path / "T.xlsx")
    wb = load_workbook(out)
    try:
        rows = _seq_summary(wb["Tracker"])
    finally:
        wb.close()
    total = next(r for r in rows if r[0] == "TOTAL")
    members, released, held, dropped = total[3], total[4], total[5], total[6]
    assert (members, released, held, dropped) == (3, 1, 1, 1)
    assert released + held + dropped == members
