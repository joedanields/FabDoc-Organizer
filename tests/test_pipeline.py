"""End-to-end and unit tests for the register pipeline."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from conftest import make_drawing

from fabdoc.categories import classify_folder_name, discover_categories, safe_sheet_name
from fabdoc.config import AppSettings, ExtractionProfile
from fabdoc.compare import compare_issues
from fabdoc.excel_out import (suggest_register_name, write_comparison_report,
                              write_register, write_validation_report)
from fabdoc.extract import (SOURCE_FILENAME, SOURCE_LARGEST, SOURCE_TITLEBLOCK,
                            extract_drawing, parse_member_mark)
from fabdoc.folder_meta import parse_folder_name
from fabdoc.memberlist import read_member_list
from fabdoc.register import build_register, sort_records
from fabdoc.register_io import read_register
from fabdoc.validate import normalise, validate


# --------------------------------------------------------------- folder names


def test_parses_title_zone_package_and_date():
    meta = parse_folder_name("Skyline Tower - Zone B - PKG-03 - 12-05-2024")
    assert meta.title == "Skyline Tower"
    assert meta.zone == "B"
    assert meta.package == "03"
    assert meta.issue_date == date(2024, 5, 12)


def test_iso_date_is_not_read_as_day_first():
    meta = parse_folder_name("Bridge Deck_Zone 4_2024-11-02")
    assert meta.issue_date == date(2024, 11, 2)
    assert meta.zone == "4"


def test_day_first_flag_switches_ambiguous_dates():
    # 05-03-2024 is 5 March day-first, 3 May month-first.
    assert parse_folder_name("Job - 05-03-2024", day_first=True).issue_date == date(2024, 3, 5)
    assert parse_folder_name("Job - 05-03-2024", day_first=False).issue_date == date(2024, 5, 3)


def test_impossible_day_first_date_falls_back():
    # 25 cannot be a month, so this must resolve as 25 December.
    assert parse_folder_name("Job - 12-25-2024").issue_date == date(2024, 12, 25)


def test_month_name_dates():
    assert parse_folder_name("Tower A - 09-Jan-2025").issue_date == date(2025, 1, 9)


def test_folder_without_metadata_becomes_the_title():
    meta = parse_folder_name("Random Package")
    assert meta.title == "Random Package"
    assert meta.issue_date is None


# ---------------------------------------------------------------- categories


@pytest.mark.parametrize("folder,expected", [
    ("Structural Drawings", "Structural"),
    ("erection", "Erection"),
    ("Part Drawings", "Part"),
    ("Single Part Drawings", "Part"),
    # An assembly drawing details one fabricated assembly for the shop; an
    # erection drawing shows where assemblies go on site. Not the same thing.
    ("Assembly Drawings", "Assembly"),
    ("Assembly", "Assembly"),
    ("Shop Drawings", "Assembly"),
    ("Random Folder", None),
])
def test_category_classification(folder, expected):
    assert classify_folder_name(folder)[0] == expected


def test_discovers_three_categories_in_order(project_folder: Path):
    cats = discover_categories(project_folder)
    assert [c.name for c in cats] == ["Structural", "Erection", "Part"]
    assert [c.count for c in cats] == [3, 2, 2]


def test_single_category_yields_one_sheet(single_category_folder: Path):
    cats = discover_categories(single_category_folder)
    assert len(cats) == 1 and cats[0].name == "Erection"


def test_unrecognised_folder_is_kept_under_its_own_name(tmp_path: Path):
    root = tmp_path / "Job X"
    make_drawing(root / "Misc Sketches" / "1_M1.pdf", "M1", seq=1)
    cats = discover_categories(root)
    assert len(cats) == 1
    assert cats[0].name == "Misc Sketches" and cats[0].is_recognised is False


def test_loose_pdfs_become_a_drawings_category(tmp_path: Path):
    root = tmp_path / "Flat Package"
    make_drawing(root / "1_A1.pdf", "A1", seq=1)
    cats = discover_categories(root)
    assert [c.name for c in cats] == ["Drawings"]


def test_sheet_names_are_deduplicated_and_clipped():
    used: set[str] = set()
    assert safe_sheet_name("Structural", used) == "Structural"
    assert safe_sheet_name("Structural", used) == "Structural (2)"
    assert len(safe_sheet_name("x" * 60)) == 31
    assert "/" not in safe_sheet_name("Part/Piece")


# ---------------------------------------------------------------- extraction


def test_extracts_labelled_fields(tmp_path: Path):
    pdf = make_drawing(tmp_path / "001_C-101.pdf", "C-101", revision=2, seq=7)
    rec = extract_drawing(pdf)
    assert rec.member_name == "C-101"
    assert rec.revision == "2"
    assert rec.seq_no == "7"
    assert rec.ok and not rec.needs_review


def test_unlabelled_mark_falls_back_to_largest_text(tmp_path: Path):
    pdf = make_drawing(tmp_path / "drawing.pdf", "B-77", labelled=False, big_mark=True)
    rec = extract_drawing(pdf)
    assert rec.member_name == "B-77"
    assert rec.member_source == SOURCE_LARGEST


def test_falls_back_to_filename_when_page_has_nothing(tmp_path: Path):
    pdf = make_drawing(tmp_path / "012_G-5_REV3.pdf", "G-5", blank=True)
    rec = extract_drawing(pdf)
    assert rec.member_name == "G-5"
    assert rec.revision == "3"
    assert rec.seq_no == "12"
    assert rec.member_source == SOURCE_FILENAME
    assert rec.needs_review  # filename-sourced rows must be flagged


def test_missing_revision_defaults_and_is_noted(tmp_path: Path):
    pdf = make_drawing(tmp_path / "no_rev.pdf", "K9", labelled=False, big_mark=True)
    rec = extract_drawing(pdf)
    assert rec.revision == "0"
    assert "revision defaulted" in rec.note_text


def test_corrupt_pdf_is_reported_not_raised(tmp_path: Path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"this is not a pdf")
    rec = extract_drawing(bad)
    assert rec.error and rec.needs_review


@pytest.mark.parametrize("label", [
    "ASSEMBLY MARK", "MEMBER NAME", "PIECE MARK", "MARK", "ASSEMBLY No",
    "SHIPPING MARK", "MEMBER ID", "PART POSITION", "ASSEMBLY REF",
])
def test_common_title_block_label_wordings(tmp_path: Path, label: str):
    """An unrecognised qualifier must not be captured as the mark.

    "MEMBER ID : B-101" previously yielded "ID", because the qualifier group
    failed to match and the capture fell onto the next word.
    """
    pdf = tmp_path / f"{label.replace(' ', '_')}.pdf"
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    page.insert_text((520, 430), f"{label} : B-101", fontsize=11)
    doc.save(pdf)
    doc.close()
    assert extract_drawing(pdf).member_name == "B-101"


@pytest.mark.parametrize("mark", [
    "B-101", "1001", "ASSY-12A", "B/101", "C101A", "EB-1", "P.301",
    "BM_204", "12", "SC-1001-A", "COL-1",
])
def test_real_world_mark_shapes(tmp_path: Path, mark: str):
    pdf = make_drawing(tmp_path / f"{mark.replace('/', '~')}.pdf", mark, revision=1, seq=1)
    assert extract_drawing(pdf).member_name == mark.upper()


def test_stopwords_are_never_taken_as_member_names(tmp_path: Path):
    pdf = make_drawing(tmp_path / "SCALE.pdf", "X1", blank=True)
    rec = extract_drawing(pdf)
    assert rec.member_name != "SCALE"


def test_bad_user_regex_does_not_crash_extraction(tmp_path: Path):
    pdf = make_drawing(tmp_path / "001_C-101.pdf", "C-101", revision=1, seq=1)
    profile = ExtractionProfile(member_patterns=["([unclosed"] + ExtractionProfile().member_patterns)
    rec = extract_drawing(pdf, profile=profile)
    assert rec.member_name == "C-101"


# ------------------------------------------------------------------ register


def test_build_register_covers_every_category(project_folder: Path):
    reg = build_register(project_folder)
    assert [c.name for c in reg.categories] == ["Structural", "Erection", "Part"]
    assert reg.total == 7
    assert reg.review_count == 0
    assert reg.meta.title == "Skyline Tower"
    assert reg.meta.zone == "B"


def test_rows_are_sorted_by_sequence_number():
    from fabdoc.extract import DrawingRecord
    records = [DrawingRecord(seq_no=s, member_name=f"M{s}") for s in ["10", "2", "1"]]
    assert [r.seq_no for r in sort_records(records)] == ["1", "2", "10"]


def test_records_sort_by_zone_then_sequence_then_mark():
    from fabdoc.extract import DrawingRecord
    records = [
        DrawingRecord(member_name="17271X9", zone="2", seq_group="271"),
        DrawingRecord(member_name="17172C10", zone="1", seq_group="172"),
        DrawingRecord(member_name="17172C2", zone="1", seq_group="172"),
        DrawingRecord(member_name="17173R1", zone="1", seq_group="173"),
    ]
    assert [r.member_name for r in sort_records(records)] == [
        "17172C2", "17172C10", "17173R1", "17271X9",
    ]


def test_records_without_a_zone_sort_last_without_error():
    from fabdoc.extract import DrawingRecord
    records = [
        DrawingRecord(member_name="LOOSE"),
        DrawingRecord(member_name="17172C1", zone="1", seq_group="172"),
    ]
    assert [r.member_name for r in sort_records(records)] == ["17172C1", "LOOSE"]


def test_category_selection_limits_the_run(project_folder: Path):
    reg = build_register(project_folder, selected_categories=["Erection"])
    assert len(reg.categories) == 1 and reg.total == 2


def test_progress_reports_every_file(project_folder: Path):
    seen: list[tuple[int, int]] = []
    build_register(project_folder, progress=lambda d, t, _l: seen.append((d, t)))
    assert seen[0] == (1, 7) and seen[-1] == (7, 7)


def test_cancellation_stops_the_run(project_folder: Path):
    reg = build_register(project_folder, should_cancel=lambda: True)
    assert reg.total == 0


# --------------------------------------------------------------------- excel


def test_workbook_has_a_sheet_per_category(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    out = write_register(reg, tmp_path / "register.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Structural", "Erection", "Part"]

    ws = wb["Structural"]
    headers = None
    for row in ws.iter_rows(values_only=True):
        if row and "Member Name" in [str(c) for c in row if c]:
            headers = [c for c in row if c]
            break
    # Title and date belong in the header band above, stated once, not
    # repeated on every row of the deliverable.
    assert headers == ["S.No", "Member Name", "Revision No"]
    wb.close()


def test_source_columns_can_be_switched_back_on(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    out = write_register(reg, tmp_path / "traced.xlsx", include_source=True)
    wb = load_workbook(out)
    ws = wb["Structural"]
    headers = next(([c for c in row if c] for row in ws.iter_rows(values_only=True)
                    if row and "Member Name" in [str(c) for c in row if c]), None)
    assert headers == ["S.No", "Member Name", "Revision No", "Source File", "Notes"]
    wb.close()


def test_single_category_workbook_has_one_register_sheet(single_category_folder: Path,
                                                         tmp_path: Path):
    reg = build_register(single_category_folder)
    out = write_register(reg, tmp_path / "one.xlsx")
    wb = load_workbook(out)
    assert [s for s in wb.sheetnames if s != "Summary"] == ["Erection"]
    wb.close()


def test_suggested_name_matches_the_source_folder(project_folder: Path):
    """The register is filed next to its issue, so the names must agree."""
    reg = build_register(project_folder)
    assert suggest_register_name(reg) == f"{project_folder.name}.xlsx"


def test_suggested_name_falls_back_when_the_folder_has_no_name(project_folder: Path):
    reg = build_register(project_folder)
    reg.project_folder = Path("D:\\")  # a drive root has no name of its own
    assert suggest_register_name(reg).endswith(".xlsx")
    assert "/" not in suggest_register_name(reg)


def test_register_round_trips_through_excel(project_folder: Path, tmp_path: Path):
    original = build_register(project_folder)
    out = write_register(original, tmp_path / "register.xlsx")
    reloaded = read_register(out)
    assert reloaded.total == original.total
    assert set(reloaded.member_names()) == set(original.member_names())
    assert reloaded.meta.title == "Skyline Tower"


# ---------------------------------------------------------------- validation


def _member_file(tmp_path: Path, names: list[str], name: str = "members.csv") -> Path:
    path = tmp_path / name
    path.write_text("Assembly Mark\n" + "\n".join(names), encoding="utf-8")
    return path


def test_clean_comparison_passes(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    members = _member_file(tmp_path, ["C-101", "C-102", "C-110", "E-201", "E-202",
                                      "P-301", "P-302"])
    result = validate(reg, read_member_list(members))
    assert result.is_clean
    assert len(result.matched) == 7
    assert result.verdict.startswith("PASS")


def test_detects_members_missing_from_drawings(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    members = _member_file(tmp_path, ["C-101", "C-102", "C-110", "E-201", "E-202",
                                      "P-301", "P-302", "C-999"])
    result = validate(reg, read_member_list(members))
    assert result.missing_in_drawings == ["C-999"]
    assert not result.is_clean


def test_detects_drawings_absent_from_the_model(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    members = _member_file(tmp_path, ["C-101", "C-102"])
    result = validate(reg, read_member_list(members))
    assert set(result.extra_in_drawings) == {"C-110", "E-201", "E-202", "P-301", "P-302"}


def test_category_filter_narrows_the_comparison(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    members = _member_file(tmp_path, ["E-201", "E-202"])
    result = validate(reg, read_member_list(members), categories=["Erection"])
    assert result.is_clean


def test_duplicate_drawings_are_flagged(tmp_path: Path):
    root = tmp_path / "Dupes - 2024-01-01"
    make_drawing(root / "Part" / "1_D1.pdf", "D1", seq=1)
    make_drawing(root / "Part" / "2_D1.pdf", "D1", seq=2)
    reg = build_register(root)
    result = validate(reg, read_member_list(_member_file(tmp_path, ["D1"])))
    assert "D1" in result.duplicates_in_drawings
    assert len(result.duplicates_in_drawings["D1"]) == 2
    assert result.is_clean  # a duplicate is a warning, not a mismatch


def test_normalisation_options():
    settings = AppSettings()
    assert normalise(" c-101 ", settings) == "C-101"
    settings.compare_strip_leading_zeros = True
    assert normalise("B007", settings) == normalise("B7", settings)


def test_case_sensitive_comparison_can_be_enforced(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    settings = AppSettings()
    settings.compare_case_insensitive = False
    members = _member_file(tmp_path, ["c-101"])
    result = validate(reg, read_member_list(members), settings=settings)
    assert "c-101" in result.missing_in_drawings


def test_validation_report_sheets(project_folder: Path, tmp_path: Path):
    reg = build_register(project_folder)
    members = _member_file(tmp_path, ["C-101", "ZZ-1"])
    result = validate(reg, read_member_list(members))
    out = write_validation_report(result, tmp_path / "report.xlsx")
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "Missing in Drawings" in wb.sheetnames
    assert "Not in Model" in wb.sheetnames
    assert wb["Missing in Drawings"]["A2"].value == "ZZ-1"
    wb.close()


# --------------------------------------------------------------- member lists


def test_reads_csv_and_picks_the_member_column(tmp_path: Path):
    path = tmp_path / "model.csv"
    path.write_text("Qty,Assembly Mark,Weight\n1,C-101,250\n1,C-102,260\n", encoding="utf-8")
    result = read_member_list(path)
    assert result.column_name == "Assembly Mark"
    assert result.members == ["C-101", "C-102"]


def test_reads_plain_text_one_per_line(tmp_path: Path):
    path = tmp_path / "model.txt"
    path.write_text("C-101\nC-102\n\nC-103\n", encoding="utf-8")
    assert read_member_list(path).members == ["C-101", "C-102", "C-103"]


def test_reads_xlsx_export(tmp_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Phase", "Assembly Position", "Weight"])
    ws.append([1, "E-201", 120])
    ws.append([1, "E-202", 130])
    path = tmp_path / "model.xlsx"
    wb.save(path)
    result = read_member_list(path)
    assert result.column_name == "Assembly Position"
    assert result.members == ["E-201", "E-202"]


def test_explicit_column_override(tmp_path: Path):
    path = tmp_path / "model.csv"
    path.write_text("Mark,Alt Mark\nA1,B1\n", encoding="utf-8")
    assert read_member_list(path, column="Alt Mark").members == ["B1"]


def test_numeric_marks_are_not_lost_to_float_formatting(tmp_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Mark"])
    ws.append([101])
    path = tmp_path / "numeric.xlsx"
    wb.save(path)
    assert read_member_list(path).members == ["101"]


def test_legacy_xls_is_rejected_with_guidance(tmp_path: Path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"")
    with pytest.raises(ValueError, match="xlsx"):
        read_member_list(path)


# ------------------------------------------------------- zones and sequences


def test_multiple_zones_are_all_captured():
    meta = parse_folder_name(
        "25. 2026-07-06  Stairs at Zone 1 and Zone 2 for Re Approval (Seqs 172,173,270,271)"
    )
    assert meta.zones == ["1", "2"]
    assert meta.zone == "1, 2"
    assert meta.issue_no == "25"
    assert meta.sequences == ["172", "173", "270", "271"]
    assert meta.issue_date == date(2026, 7, 6)


def test_zone_tokens_are_not_cut_out_of_the_title():
    """Excising the zones left a mangled "Stairs at  and Zone 2"."""
    meta = parse_folder_name(
        "25. 2026-07-06  Stairs at Zone 1 and Zone 2 for Re Approval (Seqs 172,173,270,271)"
    )
    assert meta.title == "Stairs at Zone 1 and Zone 2 for Re Approval"


def test_a_sequence_range_is_written_out():
    """Zone 1 runs "10, 11, 12, 121 thru 139 & 150" - a range, not a list."""
    meta = parse_folder_name(
        "27. 2026-08-20 Stairs at Zone 1 (Seqs 10, 11, 12, 121 thru 139 & 150)"
    )
    assert meta.sequences == ["10", "11", "12"] + [str(n) for n in range(121, 140)] + ["150"]
    assert meta.zones == ["1"]


@pytest.mark.parametrize("written", [
    "10,11,12,121-139,150",
    "10, 11, 12, 121 thru 139 & 150",
    "10 11 12, 121 through 139 and 150",
    "10,11,12,121 to 139,150",
])
def test_the_ways_a_range_is_written_all_read_alike(written):
    meta = parse_folder_name(f"27. 2026-08-20 Stairs at Zone 1 (Seqs {written})")
    assert meta.sequences[:3] == ["10", "11", "12"]
    assert meta.sequences[3:] == [str(n) for n in range(121, 140)] + ["150"]


def test_a_range_leaves_nothing_behind_in_the_title():
    """The title is the package identity the tracker chains on, so a sequence
    list that varies every issue must not survive into it."""
    first = parse_folder_name("27. 2026-08-20 Stairs at Zone 1 (Seqs 172,173)")
    later = parse_folder_name("31. 2026-09-04 Stairs at Zone 1 (Seqs 121 thru 139)")
    assert first.title == later.title == "Stairs at Zone 1"


@pytest.mark.parametrize("written", [
    "TBC",              # not numbered at all
    "139-121",          # backwards
    "12-1390",          # a typo that would expand to a thousand sequences
])
def test_brackets_that_are_not_a_sequence_list_stay_in_the_title(written):
    """Better a noisy title than sequences the package never had."""
    meta = parse_folder_name(f"27. 2026-08-20 Stairs at Zone 1 (Seqs {written})")
    assert meta.sequences == []
    assert meta.title == f"Stairs at Zone 1 (Seqs {written})"


def test_a_sequence_named_twice_is_kept_once():
    meta = parse_folder_name("27. 2026-08-20 Stairs at Zone 1 (Seqs 121-125, 123)")
    assert meta.sequences == ["121", "122", "123", "124", "125"]


@pytest.mark.parametrize("mark,job,seq,zone", [
    ("17172C172", "17", "172", "1"),
    ("17173R44", "17", "173", "1"),
    ("17270S12", "17", "270", "2"),
    ("17271X211", "17", "271", "2"),
    ("17471C10", "17", "471", "4"),
    # A package mixes sequence widths: "17120B163" is seq 120 and "1710B84" is
    # seq 10. Pinned at 3 digits the 2-digit marks got no zone at all and were
    # clustered into a second unlabelled table - 245 rows of a 968-row register.
    ("17120B163", "17", "120", "1"),
    ("1710B84", "17", "10", "1"),
    ("1710B100", "17", "10", "1"),
    ("1720S5", "17", "20", "2"),
    # The rest of Zone 1: seqs 10, 11, 12, 121-139 and 150.
    ("1711C5", "17", "11", "1"),
    ("1712A3", "17", "12", "1"),
    ("17121B4", "17", "121", "1"),
    ("17139R2", "17", "139", "1"),
    ("17150C2", "17", "150", "1"),
])
def test_sequence_and_zone_derived_from_member_mark(mark, job, seq, zone):
    assert parse_member_mark(mark) == (job, seq, zone)


def test_a_three_digit_sequence_still_wins_over_a_two_digit_reading():
    """The quantifier is greedy, so 17172C172 must not read as seq 17."""
    assert parse_member_mark("17172C172")[1] == "172"
    assert parse_member_mark("17120B163")[1] == "120"


def test_marks_that_do_not_encode_a_sequence_are_not_forced():
    assert parse_member_mark("C-101") == ("", "", "")
    assert parse_member_mark("") == ("", "", "")


def test_serial_number_restarts_within_each_sequence(tmp_path: Path):
    """S.No counts within its own band, not across the zone.

    Each band is a separate slice of work, so "the third drawing of sequence
    172" is the number worth printing. The band headings carry the counts.
    """
    root = tmp_path / "Stairs at Zone 1 and Zone 2 - 2026-07-06"
    for mark in ["17172C1", "17172C2", "17173R1", "17270S1", "17271X1"]:
        make_drawing(root / "Assembly" / f"{mark}  - Rev B.pdf", mark, revision="B")
    reg = build_register(root)
    cat = reg.categories[0]
    groups = cat.zone_groups()
    assert [z for z, _ in groups] == ["1", "2"]
    #                        seq 172   seq 173  seq 270  seq 271
    assert [r.seq_no for _, recs in groups for r in recs] == ["1", "2", "1", "1", "1"]
    assert cat.sequences_for("1") == ["172", "173"]
    assert reg.meta.zones == ["1", "2"]


def test_workbook_writes_a_banded_header_per_zone(tmp_path: Path):
    root = tmp_path / "Stairs at Zone 1 and Zone 2 - 2026-07-06"
    for mark in ["17172C1", "17271X1"]:
        make_drawing(root / "Assembly" / f"{mark}  - Rev B.pdf", mark, revision="B")
    out = write_register(build_register(root), tmp_path / "zoned.xlsx")
    wb = load_workbook(out)
    ws = wb["Assembly"]
    bands = [str(r[0]) for r in ws.iter_rows(values_only=True)
             if r and r[0] and str(r[0]).startswith("ZONE")]
    assert len(bands) == 2
    assert bands[0].startswith("ZONE 1") and "Seq 172" in bands[0]
    assert bands[1].startswith("ZONE 2") and "Seq 271" in bands[1]
    wb.close()


# ------------------------------------------------------- issue vs issue diff


def _issue(tmp_path: Path, name: str, marks: dict[str, str]) -> Path:
    root = tmp_path / name
    for mark, rev in marks.items():
        make_drawing(root / "Assembly" / f"{mark}  - Rev {rev}.pdf", mark, revision=rev)
    return root


def test_compare_issues_reports_added_and_removed(tmp_path: Path):
    old = _issue(tmp_path, "15. 2026-06-26 Stairs at Zone 1 - Approval",
                 {"17172C1": "A", "17172C2": "A", "17172C3": "A"})
    new = _issue(tmp_path, "25. 2026-07-06 Stairs at Zone 1 - Re Approval",
                 {"17172C1": "A", "17172C2": "B", "17172C9": "A"})
    result = compare_issues(build_register(old), build_register(new))
    assert [d.member_name for d in result.added] == ["17172C9"]
    assert [d.member_name for d in result.removed] == ["17172C3"]
    assert [d.member_name for d in result.revision_changed] == ["17172C2"]
    assert [d.member_name for d in result.unchanged] == ["17172C1"]
    assert not result.is_identical


def test_compare_issues_carries_old_and_new_revisions(tmp_path: Path):
    old = _issue(tmp_path, "old pkg 2026-06-26", {"17172C1": "A"})
    new = _issue(tmp_path, "new pkg 2026-07-06", {"17172C1": "B"})
    d = compare_issues(build_register(old), build_register(new)).revision_changed[0]
    assert (d.old_revision, d.new_revision) == ("A", "B")
    assert d.zone == "1"


def test_identical_issues_compare_clean(tmp_path: Path):
    old = _issue(tmp_path, "old pkg 2026-06-26", {"17172C1": "A", "17271X1": "A"})
    new = _issue(tmp_path, "new pkg 2026-07-06", {"17172C1": "A", "17271X1": "A"})
    result = compare_issues(build_register(old), build_register(new))
    assert result.is_identical
    assert result.verdict.startswith("IDENTICAL")
    assert len(result.unchanged) == 2


def test_comparison_report_sheets(tmp_path: Path):
    old = _issue(tmp_path, "old pkg 2026-06-26", {"17172C1": "A", "17172C3": "A"})
    new = _issue(tmp_path, "new pkg 2026-07-06", {"17172C1": "B", "17172C9": "A"})
    result = compare_issues(build_register(old), build_register(new))
    out = write_comparison_report(result, tmp_path / "diff.xlsx")
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "Added" in wb.sheetnames and "Removed" in wb.sheetnames
    assert wb["Added"]["A2"].value == "17172C9"
    assert wb["Removed"]["A2"].value == "17172C3"
    wb.close()


def test_a_superseded_default_pattern_is_migrated_on_load(tmp_path: Path):
    """A fix in the shipped defaults must reach engineers who saved settings.

    A saved profile holds a copy of every pattern, so the 3-digit sequence
    pattern - a default nobody chose - kept shadowing its own fix. On the real
    968-drawing package that silently dropped 245 rows out of their sequences.
    """
    import json
    from fabdoc.config import ExtractionProfile, load_settings, save_settings, AppSettings

    old = r"^(?P<job>\d{2})(?P<seq>\d{3})(?P<rest>[A-Za-z].*)$"
    settings = AppSettings()
    settings.profile.member_seq_pattern = old
    path = tmp_path / "settings.json"
    save_settings(settings, path)
    assert json.loads(path.read_text(encoding="utf-8"))["profile"]["member_seq_pattern"] == old

    loaded = load_settings(path)
    assert loaded.profile.member_seq_pattern == ExtractionProfile().member_seq_pattern

    from fabdoc.extract import parse_member_mark
    assert parse_member_mark("1710B84", loaded.profile)[1] == "10"


def test_a_pattern_the_engineer_tuned_is_never_overwritten(tmp_path: Path):
    """Migration only replaces values identical to a superseded default."""
    from fabdoc.config import load_settings, save_settings, AppSettings

    mine = r"^(?P<job>\d{3})(?P<seq>\d{2})(?P<rest>[A-Za-z].*)$"
    settings = AppSettings()
    settings.profile.member_seq_pattern = mine
    path = tmp_path / "settings.json"
    save_settings(settings, path)
    assert load_settings(path).profile.member_seq_pattern == mine


# ------------------------------------------------- single-part title blocks


def _part_titleblock(path: Path, mark: str, material: str,
                     length: str = "0'-11\"") -> Path:
    """A single-part title block: column headings with values on another row.

    Nothing here is on the same line as its heading, which is what breaks a
    same-line label pattern. The mark is the largest text in the block, exactly
    as the real drawings present it.
    """
    doc = fitz.open()
    page = doc.new_page(width=792, height=612)
    page.insert_text((672, 568), "Part #", fontsize=9)
    page.insert_text((672, 578), mark, fontsize=13.5)      # the real mark
    page.insert_text((671, 481), "Qty", fontsize=9)
    page.insert_text((715, 481), "In Assembly", fontsize=9)
    page.insert_text((671, 465), "1", fontsize=9)
    page.insert_text((715, 465), material, fontsize=9)     # the trap
    page.insert_text((533, 552), "Material", fontsize=9)
    page.insert_text((625, 552), "Length", fontsize=9)
    page.insert_text((625, 565), length, fontsize=9)
    doc.save(path)
    doc.close()
    return path


@pytest.mark.parametrize("material", [
    "HSS4X4X1/2",        # a section size
    "C12X25",
    "PIPE1-1/2SCH40",
    "PL1/2",
    "15/16",             # a fraction off a dimension
    "29.00",             # a length
    "CENTERLINE",        # a note with no digits at all
])
def test_a_part_mark_is_not_confused_with_its_material(tmp_path: Path, material: str):
    """"In Assembly" is a column heading, not a label introducing a value.

    It matched the ASSEMBLY pattern and the capture fell onto the cell beside
    it, so 121 of 125 single-part drawings came out named after a steel section
    or a length - and none was flagged, because a labelled capture counts as
    confident.
    """
    pdf = _part_titleblock(tmp_path / f"{material.replace('/', '~')}.pdf",
                           "17172C250", material)
    rec = extract_drawing(pdf, category="Part")
    assert rec.member_name == "17172C250", f"captured {rec.member_name!r} instead"


def test_a_genuine_assembly_label_still_wins(tmp_path: Path):
    """The fix must not cost the templates that do label their marks."""
    pdf = make_drawing(tmp_path / "labelled.pdf", "C-101", revision=2, seq=7)
    rec = extract_drawing(pdf)
    assert rec.member_name == "C-101"
    assert rec.member_source == SOURCE_TITLEBLOCK


@pytest.mark.parametrize("wording", ["ASSEMBLY MARK", "ASSEMBLY No", "ASSEMBLY REF",
                                     "MEMBER NAME", "MEMBER ID"])
def test_labels_still_match_with_a_qualifier_or_separator(tmp_path: Path, wording: str):
    pdf = tmp_path / f"{wording.replace(' ', '_')}.pdf"
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    page.insert_text((520, 430), f"{wording} : B-101", fontsize=11)
    doc.save(pdf)
    doc.close()
    assert extract_drawing(pdf).member_name == "B-101"


def test_superseded_label_patterns_are_migrated_too(tmp_path: Path):
    """The saved profile shadowed the label fix as well as the sequence one.

    A profile saved before the fix carries the loose patterns, the short
    stopword list, and no reject list at all - so a settings file written once,
    months ago, kept every single-part drawing named after its material.
    """
    import json
    from fabdoc.config import (SUPERSEDED_DEFAULTS, ExtractionProfile, AppSettings,
                               load_settings, save_settings)

    settings = AppSettings()
    settings.profile.member_patterns = list(SUPERSEDED_DEFAULTS["member_patterns"][0])
    settings.profile.member_stopwords = list(SUPERSEDED_DEFAULTS["member_stopwords"][0])
    path = tmp_path / "settings.json"
    save_settings(settings, path)

    # A field added after the file was written is simply absent from it.
    raw = json.loads(path.read_text(encoding="utf-8"))
    del raw["profile"]["member_reject_patterns"]
    path.write_text(json.dumps(raw), encoding="utf-8")

    fresh = ExtractionProfile()
    loaded = load_settings(path).profile
    assert loaded.member_patterns == fresh.member_patterns
    assert loaded.member_stopwords == fresh.member_stopwords
    assert loaded.member_reject_patterns == fresh.member_reject_patterns


def test_a_label_pattern_the_engineer_wrote_is_kept(tmp_path: Path):
    from fabdoc.config import AppSettings, load_settings, save_settings

    mine = [r"MY\s*MARK\s*[:\-]\s*([A-Z0-9\-]+)"]
    settings = AppSettings()
    settings.profile.member_patterns = mine
    path = tmp_path / "settings.json"
    save_settings(settings, path)
    assert load_settings(path).profile.member_patterns == mine


# ------------------------------------------------------------- mark casing


@pytest.mark.parametrize("mark", ["17ch104", "17hsp1", "17a24", "17172C172", "17CH104"])
def test_the_mark_is_recorded_as_the_drawing_carries_it(tmp_path: Path, mark: str):
    """A mark is an identifier, not a heading.

    Single-part drawings are marked in lower case and assemblies in upper, in
    the same package. Folding everything to upper case made the register
    disagree with the drawing it came from.
    """
    pdf = make_drawing(tmp_path / f"{mark}.pdf", mark, revision=0, big_mark=True,
                       labelled=False)
    assert extract_drawing(pdf).member_name == mark


def test_case_is_preserved_from_a_label_and_from_the_filename(tmp_path: Path):
    labelled = make_drawing(tmp_path / "labelled.pdf", "17ch55", revision="A")
    assert extract_drawing(labelled).member_name == "17ch55"

    blank = make_drawing(tmp_path / "17hsp274  - Rev 0.pdf", "x", blank=True)
    from_name = extract_drawing(blank)
    assert from_name.member_name == "17hsp274"
    assert from_name.member_source == SOURCE_FILENAME


def test_a_revision_is_still_upper_cased(tmp_path: Path):
    """Revisions are a sequence, not an identity - "a" and "A" are one rev."""
    pdf = make_drawing(tmp_path / "r.pdf", "17ch55", revision="b")
    assert extract_drawing(pdf).revision == "B"


def test_mixed_case_marks_do_not_split_a_band(tmp_path: Path):
    """Grouping is case-insensitive even though display is not."""
    root = tmp_path / "Parts - 2026-07-22"
    for mark in ["17ch1", "17CH2", "17Ch3"]:
        make_drawing(root / "Single Part Drawings" / f"{mark}.pdf", mark,
                     revision=0, big_mark=True, labelled=False)
    cat = build_register(root).categories[0]
    bands = cat.band_groups("")
    assert [b for b, _ in bands] == [("type", "CH")]
    assert len(bands[0][1]) == 3
    assert sorted(r.member_name for r in bands[0][1]) == ["17CH2", "17Ch3", "17ch1"]


# ------------------------------------------------- quantity and cut length


def test_quantity_and_length_are_read_from_the_title_block_table(tmp_path: Path):
    """The two numbers the shop works to, read by column rather than by label.

    They are written as a table - the heading in one row, the value in the row
    below - so read as flowing text "Qty" is followed by "Profile" and no
    same-line pattern can reach the number.
    """
    pdf = make_drawing(tmp_path / "part.pdf", "17a25", revision=0,
                       qty="3", length="3'-11 5/8\"")
    rec = extract_drawing(pdf, category="Part")
    assert rec.quantity == "3"
    assert rec.length == "3'-11 5/8\""
    assert rec.length_inches == 47.625
    assert rec.quantity_source == SOURCE_TITLEBLOCK
    assert rec.length_source == SOURCE_TITLEBLOCK


def test_a_drawing_without_them_says_nothing_rather_than_guessing(tmp_path: Path):
    """An approval sheet carries no quantity, and a wrong one is worse than none."""
    pdf = make_drawing(tmp_path / "approval.pdf", "17172C172", revision="A")
    rec = extract_drawing(pdf, category="Part")
    assert rec.quantity == "" and rec.length == ""
    assert rec.length_inches is None


def test_only_single_part_drawings_are_read_for_them(tmp_path: Path):
    """An assembly title block carries a count too - of assemblies, not of cuts.

    Tracked in the same column as a part's quantity, the two read as one number
    meaning two different things.
    """
    pdf = make_drawing(tmp_path / "assembly.pdf", "17172C172", revision=0,
                       qty="3", length="3'-11 5/8\"")
    assert extract_drawing(pdf, category="Assembly").quantity == ""
    assert extract_drawing(pdf, category="Part").quantity == "3"


@pytest.mark.parametrize("stem, drawn", [("17HSP134", "17hsp134"),
                                         ("17hsp135", "17HSP135")])
def test_the_file_name_decides_the_case_of_the_mark(tmp_path: Path, stem, drawn):
    """The register is worked from on a shop floor: it says what the file says."""
    pdf = make_drawing(tmp_path / f"{stem}  - Rev 0.pdf", drawn, revision=0)
    assert extract_drawing(pdf, category="Part").member_name == stem


def test_a_filename_naming_a_different_mark_does_not_override_it(tmp_path: Path):
    """Only the spelling is taken, and only when the two are the same mark."""
    pdf = make_drawing(tmp_path / "17hsp999  - Rev 0.pdf", "17HSP134", revision=0)
    assert extract_drawing(pdf, category="Part").member_name == "17HSP134"


def test_the_register_shows_them_only_when_the_drawings_carry_them(tmp_path: Path):
    root = tmp_path / "Zone 1 for Fabrication"
    for i, (mark, qty, length) in enumerate(
            [("17a25", "3", "3'-11 5/8\""), ("17a26", "1", "2'-0\"")], start=1):
        make_drawing(root / "Single Part Drawings" / f"{i}_{mark}.pdf", mark,
                     revision=0, seq=i, qty=qty, length=length)

    out = write_register(build_register(root), tmp_path / "R.xlsx")
    wb = load_workbook(out)
    try:
        ws = wb[[s for s in wb.sheetnames if s != "Summary"][0]]
        rows = [[c for c in row if c is not None]
                for row in ws.iter_rows(values_only=True)]
        headers = next(r for r in rows if "Member Name" in [str(c) for c in r])
        assert headers == ["S.No", "Member Name", "Revision No", "Qty", "Length"]
        row = next(r for r in rows if "17a25" in [str(c) for c in r])
        assert row[3] == 3 and row[4] == "3'-11 5/8\""
    finally:
        wb.close()


def test_a_register_read_back_still_carries_them(tmp_path: Path):
    """The comparison and validation screens work off a written register."""
    root = tmp_path / "Zone 1 for Fabrication"
    make_drawing(root / "Single Part Drawings" / "1_17a25.pdf", "17a25",
                 revision=0, seq=1, qty="3", length="3'-11 5/8\"")
    out = write_register(build_register(root), tmp_path / "R.xlsx")

    back = read_register(out)
    rec = back.all_records()[0]
    assert rec.quantity == "3"
    assert rec.length == "3'-11 5/8\"" and rec.length_inches == 47.625
